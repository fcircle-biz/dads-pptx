# -*- coding: utf-8 -*-
"""DADS準拠PPTXの最小構成サンプル（表紙・カード・表・出典の4枚）。

実行例:
    .venv/bin/python .agents/skills/dads-pptx/examples/build_example.py /tmp/example.pptx
"""
import sys, os

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))

from pptx import Presentation
from dads_theme import *   # noqa

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "example.pptx")
DECK_TITLE = "サンプル資料"
CONTENT_W = CANVAS_W - MARGIN_X * 2

prs = Presentation()
prs.slide_width, prs.slide_height = px(CANVAS_W), px(CANVAS_H)
BLANK = prs.slide_layouts[6]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, CANVAS_W, CANVAS_H, fill=bg)
    return s


def header(slide, kicker, title, lead=None):
    rect(slide, MARGIN_X, 52, 6, 26, fill=PRIMARY)
    block(slide, MARGIN_X + 16, 52, CONTENT_W - 16, 26,
          [dict(text=kicker, style="Oln-14B-100", color=PRIMARY)], anchor="m")
    block(slide, MARGIN_X, 84, CONTENT_W, 48,
          [dict(text=title, style="Std-32B-150", color=INK)])
    rule(slide, MARGIN_X, 140, CONTENT_W, 2, BORDER)
    if lead:
        block(slide, MARGIN_X, 152, CONTENT_W, 30,
              [dict(text=lead, style="Std-17N-170", color=INK_SUB)])


def footer(slide, no):
    rule(slide, MARGIN_X, 660, CONTENT_W, 1, BORDER_SOFT)
    block(slide, MARGIN_X, 672, 700, 20,
          [dict(text=DECK_TITLE, style="Dns-14N-130", color=INK_MUTE)])
    block(slide, CANVAS_W - MARGIN_X - 200, 672, 200, 20,
          [dict(text=str(no), style="Dns-14B-130", color=INK_MUTE, align="r")])


def datatable(slide, x, y, w, cols, rows, row_h=56, head_h=48):
    widths = [w * c for c in cols["ratio"]]
    xs, acc = [], x
    for cw in widths:
        xs.append(acc)
        acc += cw
    rect(slide, x, y, w, head_h, fill=BG_KEY)
    for i, label in enumerate(cols["labels"]):
        block(slide, xs[i] + 16, y, widths[i] - 32, head_h,
              [dict(text=label, style="Dns-14B-130", color=PRIMARY)], anchor="m")
    rule(slide, x, y + head_h, w, 2, PRIMARY)
    cy = y + head_h + 2
    for r in rows:
        for i, cell in enumerate(r):
            block(slide, xs[i] + 16, cy, widths[i] - 32, row_h,
                  [dict(text=cell,
                        style=("Std-16B-170" if i == 0 else "Std-16N-170"),
                        color=(INK if i == 0 else INK_SUB))], anchor="m")
        rule(slide, x, cy + row_h, w, 1, BORDER_SOFT)
        cy += row_h + 1
    return cy


# ---------------------------------------------------------------- 表紙
s = new_slide(TERTIARY)
rect(s, 0, 0, 8, CANVAS_H, fill=SECONDARY)
block(s, MARGIN_X + 24, 200, 900, 40,
      [dict(text="SUBTITLE IN LATIN", style="Oln-16B-100", color=BLUE[300])])
block(s, MARGIN_X + 24, 246, 1000, 90,
      [dict(text="資料タイトル", style="Dsp-64B-140", color=WHITE)])
rect(s, MARGIN_X + 24, 372, 120, 4, fill=SECONDARY)
block(s, MARGIN_X + 24, 404, 860, 40,
      [dict(text="サブタイトルを1〜2行で置く", style="Std-24N-150", color=BLUE[100])])
rule(s, MARGIN_X + 24, 600, CONTENT_W - 48, 1, BLUE[1100])
block(s, MARGIN_X + 24, 620, 600, 24,
      [dict(text="組織名", style="Dns-16N-130", color=BLUE[200])])

# ---------------------------------------------------------------- カード
s = new_slide()
header(s, "01 ── セクション名", "カードを3枚並べる型", "1枚のスライドで伝えることは1つに絞る。")
cards = [("見出しを短く", "説明文は2〜3行に収める。長い場合はスライドを分ける。"),
         ("並列の情報に使う", "対等な要素を並べるときの型。順序があるならフローにする。"),
         ("枠線は3:1以上", "Gray-420を使う。Gray-300は非テキスト要件を満たさない。")]
cw = (CONTENT_W - 32) / 3
for i, (t1, t2) in enumerate(cards):
    x = MARGIN_X + i * (cw + 16)
    rect(s, x, 200, cw, 200, fill=SURFACE, line=BORDER, radius=RADIUS["none"])
    rect(s, x, 200, cw, 6, fill=PRIMARY)
    block(s, x + 24, 232, cw - 48, 150,
          [dict(text=t1, style="Std-20B-150", color=INK),
           dict(text=t2, style="Std-16N-170", color=INK_SUB, space_before=8)])
rect(s, MARGIN_X, 432, CONTENT_W, 72, fill=BG_KEY, line=None)
block(s, MARGIN_X + 32, 432, CONTENT_W - 64, 72,
      [dict(text="注記や補足はバックグラウンドカラーのブロックに置く。",
            style="Std-17N-170", color=TERTIARY)], anchor="m")
block(s, MARGIN_X, 536, CONTENT_W, 90,
      [dict(text="箇条書きは bullet を指定するとぶら下げインデントが効く。",
            style="Std-17N-170", color=INK, bullet="・"),
       dict(text="折り返した2行目も行頭が揃うため、長い文でも崩れない。",
            style="Std-17N-170", color=INK, bullet="・")])
footer(s, 2)

# ---------------------------------------------------------------- 表
s = new_slide()
header(s, "02 ── セクション名", "表は横罫線で組む")
datatable(s, MARGIN_X, 200, CONTENT_W,
          {"labels": ["観点", "内容", "備考"], "ratio": [0.22, 0.40, 0.38]},
          [["1行目", "本文は Std-16N-170", "先頭列だけ太字にする"],
           ["2行目", "行の高さは56が標準", "セル内は上下中央"],
           ["3行目", "ヘッダーは Blue-50 地", "下端に2pxのキーカラー罫線"]])
footer(s, 3)

# ---------------------------------------------------------------- 出典
s = new_slide()
header(s, "APPENDIX", "本資料について（出典・注記）")
notes = [("デザイン", "デジタル庁デザインシステムβ版の基本デザインを参考に構成しています。"),
         ("カラー", "テキストは4.5:1以上、非テキスト要素は3:1以上のコントラスト比を確保しています。"),
         ("書体", "Noto Sans JP（SIL Open Font License 1.1）。")]
y = 200
for t1, t2 in notes:
    badge(s, MARGIN_X, y + 4, t1, fill=BG_KEY, fg=PRIMARY, w=104)
    block(s, MARGIN_X + 128, y, CONTENT_W - 128, 40,
          [dict(text=t2, style="Std-16N-170", color=INK)])
    y += 72
rule(s, MARGIN_X, 470, CONTENT_W, 2, BORDER)
block(s, MARGIN_X, 492, CONTENT_W, 60,
      [dict(text="出典：デジタル庁デザインシステムウェブサイト https://design.digital.go.jp/dads/ のコンテンツを加工して作成",
            style="Std-16N-170", color=INK_SUB),
       dict(text="本資料はデジタル庁が作成・監修したものではありません。",
            style="Dns-14N-130", color=INK_MUTE)])
footer(s, 4)

prs.save(OUT)
print("saved:", OUT)
