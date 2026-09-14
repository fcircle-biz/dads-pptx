# -*- coding: utf-8 -*-
"""DADS準拠PPTXにアニメーション（フェードのみ・クリック進行）を付けるための共通モジュール。

python-pptx にはアニメーションのAPIがないため、スライドXMLへ直接
<p:transition>（画面切り替え）と <p:timing>（オブジェクトのアニメーション）を書き込む。

方針（references/animation.md）:
- 効果はフェードのみ。飛び込み・回転・点滅・退場は使わない
- クリック進行を既定とし、全要素が表示された最終状態だけで内容が伝わるようにする

使い方:
    from dads_anim import grab, Timeline

    tl = Timeline(slide)
    with grab(slide) as g:
        rect(slide, ...); block(slide, ...)
    tl.click("カードを表示", g.shapes)   # クリックで表示
    tl.then(other_shapes)                # 直前の表示に続けて自動で表示
    tl.apply()                           # XMLへ書き込む（アニメーションなしのスライドも apply() で切り替えを付ける）

検証:
    python dads_anim.py out/deck.pptx     # クリック数・対象図形の実在・効果の種類を検査（問題があれば終了コード1）
"""
import itertools
import re
import sys

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from dads_theme import (INK, PRIMARY, SECONDARY, BG_KEY, GRAY, STYLE, contrast,
                        rect, textbox, _apply_font)

FADE_MS = 400            # 1要素のフェード時間
THEN_GAP_MS = 120        # then() の既定の間隔
MAX_CLICKS = 6           # これを超えるスライドは分割を検討する（check で警告）
TRANSITION_SPEED = "med"


# ---------------------------------------------------------------- 図形の収集
class grab:
    """with ブロック内でスライドに追加した図形を集める（アニメーション対象の指定用）。

    グループ図形の中に追加した図形は対象外。ヘルパは slide.shapes に直接追加する前提。
    """

    def __init__(self, slide):
        self.slide = slide
        self.shapes = []

    def __enter__(self):
        self.n = len(self.slide.shapes)
        return self

    def __exit__(self, *exc):
        self.shapes = list(self.slide.shapes)[self.n:]


# ---------------------------------------------------------------- 画面切り替え
def add_transition(slide, speed=TRANSITION_SPEED):
    """画面切り替え：フェード。speed は "fast" / "med" / "slow"。"""
    sld = slide._element
    if sld.find(qn("p:transition")) is None:
        sld.find(qn("p:clrMapOvr")).addnext(
            parse_xml(f'<p:transition {nsdecls("p")} spd="{speed}"><p:fade/></p:transition>'))


# ---------------------------------------------------------------- クリック進行のフェード
class Timeline:
    """クリックごとのフェード表示を組み立てる。

    click(label, shapes)  … クリックで開始するグループ
    then(shapes, gap)     … 直前の表示が終わってから自動で続く（同じクリック内）
    apply(notes=True)     … XMLへ書き込み、発表者ノートにクリック手順を追記する
    """

    def __init__(self, slide):
        self.slide = slide
        self.groups = []

    @property
    def labels(self):
        return [g["label"] for g in self.groups]

    def click(self, label, shapes):
        self.groups.append(dict(label=label, subs=[dict(start=0, shapes=list(shapes))]))
        return self

    def then(self, shapes, gap=THEN_GAP_MS):
        assert self.groups, "then() の前に click() が必要"
        subs = self.groups[-1]["subs"]
        subs.append(dict(start=subs[-1]["start"] + FADE_MS + gap, shapes=list(shapes)))
        return self

    def apply(self, notes=True, transition=True):
        if transition:
            add_transition(self.slide)
        if not self.groups:
            return self
        sld = self.slide._element
        assert sld.find(qn("p:timing")) is None, "apply() は1スライドに1回だけ"
        ids = itertools.count(3)
        clicks, spids = [], []
        for g in self.groups:
            outer = next(ids)
            inner_xml = []
            for si, sub in enumerate(g["subs"]):
                assert sub["shapes"], f"表示する図形がない: {g['label']}"
                inner = next(ids)
                effects = []
                for k, shp in enumerate(sub["shapes"]):
                    spid = shp.shape_id
                    spids.append(spid)
                    node = ("clickEffect" if si == 0 else "afterEffect") if k == 0 else "withEffect"
                    effects.append(_fade_in(next(ids), next(ids), next(ids), spid, node))
                inner_xml.append(
                    f'<p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst>'
                    f'<p:cond delay="{sub["start"]}"/></p:stCondLst><p:childTnLst>'
                    + "".join(effects) + '</p:childTnLst></p:cTn></p:par>')
            clicks.append(
                f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst>'
                f'<p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                + "".join(inner_xml) + '</p:childTnLst></p:cTn></p:par>')
        assert len(spids) == len(set(spids)), "同じ図形を二重に登場させている"
        bld = "".join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in spids)
        timing = parse_xml(
            f'<p:timing {nsdecls("p")}><p:tnLst><p:par>'
            f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            f'<p:childTnLst>' + "".join(clicks) + '</p:childTnLst></p:cTn>'
            f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
            f'</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/>'
            f'</p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
            f'<p:bldLst>{bld}</p:bldLst></p:timing>')
        ext = sld.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)
        if notes:
            tf = self.slide.notes_slide.notes_text_frame
            steps = "クリック手順\n" + "\n".join(f"{i}. {l}" for i, l in enumerate(self.labels, 1))
            tf.text = (tf.text + "\n\n" + steps) if tf.text else steps
        return self


def _fade_in(eid, set_id, anim_id, spid, node):
    return (
        f'<p:par><p:cTn id="{eid}" presetID="10" presetClass="entr" presetSubtype="0" '
        f'fill="hold" grpId="0" nodeType="{node}">'
        f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:set><p:cBhvr><p:cTn id="{set_id}" dur="1" fill="hold"><p:stCondLst>'
        f'<p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/>'
        f'</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>'
        f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
        f'<p:animEffect transition="in" filter="fade"><p:cBhvr>'
        f'<p:cTn id="{anim_id}" dur="{FADE_MS}"/><p:tgtEl><p:spTgt spid="{spid}"/>'
        f'</p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')


# ---------------------------------------------------------------- 段階表示用の部品：コード＋行ハイライト
CODE_STYLE = "Std-18N-160"
CODE_BG = GRAY[50]
CODE_KEYWORDS = re.compile(
    r"(\bSELECT\b|\bFROM\b|\bWHERE\b|\bJOIN\b|\bON\b|\bGROUP BY\b|\bORDER BY\b|\bHAVING\b|"
    r"\bCASE\b|\bWHEN\b|\bTHEN\b|\bELSE\b|\bEND\b|\bAS\b|\bAND\b|\bOR\b|\bNOT\b|\bNULL\b|"
    r"\bSUM\b|\bCOUNT\b|\bAVG\b|\bMIN\b|\bMAX\b|\bCOALESCE\b|\bFILTER\b|\bPIVOT\b|\bFOR\b|\bIN\b)")


def code_line(slide, x, y, w, h, text, bg=CODE_BG, keywords=CODE_KEYWORDS):
    """コード1行を1テキストボックスで置く。keywords に一致した語をプライマリー色の太字にする。

    行ごとに独立させるのは、レンダラによる行間の差で行位置とハイライトがずれないようにするため。
    書体は Noto Sans JP（Noto Sans Mono は未導入環境で別書体に置き換わるため使わない）。
    """
    size_px, _, _ = STYLE[CODE_STYLE]
    assert contrast(INK, bg) >= 4.5 and contrast(PRIMARY, bg) >= 4.5
    tb = textbox(slide, x, y, w, h, anchor="m")
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.line_spacing = 1.3
    for part in (keywords.split(text) if keywords else [text]):
        if not part:
            continue
        run = p.add_run()
        run.text = part
        kw = bool(keywords and keywords.fullmatch(part))
        _apply_font(run, size_px, kw, PRIMARY if kw else INK)
    return tb


def code_block(slide, x, y, w, lines, lh=32, pad=20, highlights=()):
    """コード面・行・ハイライト帯を置き、dict(lines, highlights, bottom) を返す。

    highlights: [(開始行, 終了行), ...]（0始まり・両端含む）。帯は文字より背面に作られるので、
    返り値の highlights[i]（図形リスト）を Timeline.click() に渡すと、その行が順に強調される。
    コード行 lines[i] も個別に Timeline へ渡せる。
    """
    h = pad * 2 + lh * len(lines)
    rect(slide, x, y, w, h, fill=CODE_BG)
    rect(slide, x, y, 4, h, fill=PRIMARY)
    hl = []
    for a, b in highlights:
        with grab(slide) as g:
            rect(slide, x + 4, y + pad + a * lh, w - 4, lh * (b - a + 1), fill=BG_KEY)
            rect(slide, x + 4, y + pad + a * lh, 4, lh * (b - a + 1), fill=SECONDARY)
        hl.append(g.shapes)
    rows = [code_line(slide, x + 28, y + pad + i * lh, w - 44, lh, ln)
            for i, ln in enumerate(lines)]
    return dict(lines=rows, highlights=hl, bottom=y + h)


# ---------------------------------------------------------------- 検証
def check(path):
    """PPTXのアニメーションを検査する。(スライドごとの要約, エラー, 警告) を返す。"""
    from pptx import Presentation
    prs = Presentation(path)
    summary, errors, warns = [], [], []
    for no, slide in enumerate(prs.slides, 1):
        sld = slide._element
        ids = {int(e.get("id")) for e in sld.iter(qn("p:cNvPr"))}
        if sld.find(qn("p:transition")) is None:
            warns.append(f"slide{no:02d} 画面切り替えがない")
        timing = sld.find(qn("p:timing"))
        clicks = effects = 0
        if timing is not None:
            main = next((c for c in timing.iter(qn("p:cTn")) if c.get("nodeType") == "mainSeq"), None)
            if main is not None:
                clicks = sum(1 for par in main.find(qn("p:childTnLst"))
                             if par.find(qn("p:cTn")).find(qn("p:stCondLst"))
                             .find(qn("p:cond")).get("delay") == "indefinite")
            seen = set()
            for ctn in timing.iter(qn("p:cTn")):
                cls = ctn.get("presetClass")
                if cls is None:
                    continue
                effects += 1
                filt = [a.get("filter") for a in ctn.iter(qn("p:animEffect"))]
                if cls != "entr" or ctn.get("presetID") != "10" or filt != ["fade"]:
                    errors.append(f"slide{no:02d} フェード以外の効果がある（presetClass={cls}, "
                                  f"presetID={ctn.get('presetID')}）")
            for tgt in timing.iter(qn("p:spTgt")):
                spid = int(tgt.get("spid"))
                if spid not in ids:
                    errors.append(f"slide{no:02d} 存在しない図形 id={spid} を参照している")
            for ctn in timing.iter(qn("p:cTn")):
                if ctn.get("presetClass"):
                    spid = next(ctn.iter(qn("p:spTgt"))).get("spid")
                    if spid in seen:
                        errors.append(f"slide{no:02d} 図形 id={spid} を二重に登場させている")
                    seen.add(spid)
        if clicks > MAX_CLICKS:
            warns.append(f"slide{no:02d} クリックが{clicks}回（{MAX_CLICKS}回以下を目安にスライドを分ける）")
        summary.append((no, clicks, effects))
    return summary, errors, warns


if __name__ == "__main__":
    summary, errors, warns = check(sys.argv[1])
    for no, clicks, effects in summary:
        print(f"slide{no:02d}  clicks={clicks}  effects={effects}")
    for w in warns:
        print("WARN", w)
    for e in errors:
        print("ERROR", e)
    print(f"checked {len(summary)} slides, {sum(c for _, c, _ in summary)} clicks, "
          f"{len(warns)} warnings, {len(errors)} errors")
    sys.exit(1 if errors else 0)
