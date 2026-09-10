# -*- coding: utf-8 -*-
"""生成したPPTXをPillowで近似レンダリングし、はみ出し・重なりを目視確認する。

PowerPointのレンダラではないため厳密ではないが、テキストの折り返しと図形の位置関係を
同じ規則（1280x720 CSS px キャンバス、font-size px = pt/0.75）で再現して検証する。
"""
import sys, os
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

SRC = sys.argv[1]
OUTDIR = sys.argv[2]
SCALE = 1.5
EMU_PX = 914400 / 96
FONT_R = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FONT_B = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
_cache = {}


def font(size_px, bold):
    key = (round(size_px), bold)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONT_B if bold else FONT_R,
                                         int(round(size_px * SCALE)))
    return _cache[key]


def e2p(v):
    return (v or 0) / EMU_PX * SCALE


def rgb_of(color_obj, default=None):
    try:
        return "#" + str(color_obj.rgb)
    except Exception:
        return default


def wrap(text, fnt, max_w, draw):
    """CJK対応の折り返し（禁則は簡略化）。"""
    lines, cur = [], ""
    for ch in text:
        t = cur + ch
        if draw.textlength(t, font=fnt) > max_w and cur:
            lines.append(cur)
            cur = ch
        else:
            cur = t
    lines.append(cur)
    return lines


def render(slide, idx):
    W, H = int(1280 * SCALE), int(720 * SCALE)
    img = Image.new("RGB", (W, H), "#FFFFFF")
    d = ImageDraw.Draw(img)
    overflow = []
    for shp in slide.shapes:
        x, y, w, h = e2p(shp.left), e2p(shp.top), e2p(shp.width), e2p(shp.height)
        # --- 図形の塗り・線
        if shp.shape_type is not None and shp.has_text_frame and shp.shape_type != 17:
            pass
        try:
            fill = rgb_of(shp.fill.fore_color) if shp.fill.type == 1 else None
        except Exception:
            fill = None
        try:
            line = rgb_of(shp.line.color) if shp.line.fill.type == 1 else None
        except Exception:
            line = None
        if fill or line:
            prst = ""
            try:
                g = shp._element.spPr.find(qn("a:prstGeom"))
                prst = g.get("prst") if g is not None else ""
            except Exception:
                pass
            name = prst or str(shp.shape_type)
            box = [x, y, x + w, y + h]
            if name == "ellipse":
                d.ellipse(box, fill=fill, outline=line, width=max(1, int(SCALE)))
            elif name == "triangle":
                cx, cy = (x + x + w) / 2, (y + y + h) / 2
                rot = getattr(shp, "rotation", 0)
                if abs(rot - 90) < 1:      # 右向き
                    pts = [(x, y), (x + w, cy), (x, y + h)]
                else:
                    pts = [(cx, y), (x + w, y + h), (x, y + h)]
                d.polygon(pts, fill=fill or line)
            elif name == "roundRect":
                adj = 0.0
                try:
                    adj = shp.adjustments[0]
                except Exception:
                    pass
                r = adj * min(w, h)
                d.rounded_rectangle(box, radius=r, fill=fill, outline=line,
                                    width=max(1, int(SCALE)))
            else:
                d.rectangle(box, fill=fill, outline=line, width=max(1, int(SCALE)))
        # --- テキスト
        if not shp.has_text_frame:
            continue
        tf = shp.text_frame
        ml, mr = e2p(tf.margin_left), e2p(tf.margin_right)
        mt, mb = e2p(tf.margin_top), e2p(tf.margin_bottom)
        tx, tw = x + ml, w - ml - mr
        rows = []
        for p in tf.paragraphs:
            txt = "".join(r.text for r in p.runs)
            if not txt:
                continue
            r0 = p.runs[0]
            size_px = (r0.font.size.pt if r0.font.size else 18) / 0.75
            bold = bool(r0.font.bold)
            col = rgb_of(r0.font.color, "#000000")
            fnt = font(size_px, bold)
            lh = (p.line_spacing or 1.2) * size_px * SCALE
            sb = (p.space_before.pt / 0.75 * SCALE) if p.space_before else 0
            sa = (p.space_after.pt / 0.75 * SCALE) if p.space_after else 0
            pPr = p._p.find(qn("a:pPr"))
            bu, marL = None, 0
            if pPr is not None:
                bc = pPr.find(qn("a:buChar"))
                if bc is not None:
                    bu = bc.get("char")
                    marL = float(pPr.get("marL", 0)) / EMU_PX * SCALE
            avail = tw - marL
            wrapped = wrap(txt, fnt, avail, d)
            rows.append((wrapped, fnt, col, lh, sb, sa, p.alignment, bu, marL))
        total = sum(len(r[0]) * r[3] + r[4] + r[5] for r in rows)
        if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
            cy = y + mt + (h - mt - mb - total) / 2
        elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
            cy = y + h - mb - total
        else:
            cy = y + mt
        start_y = cy
        for wrapped, fnt, col, lh, sb, sa, align, bu, marL in rows:
            cy += sb
            for i, ln in enumerate(wrapped):
                lw = d.textlength(ln, font=fnt)
                if align == PP_ALIGN.CENTER:
                    lx = tx + (tw - lw) / 2
                elif align == PP_ALIGN.RIGHT:
                    lx = tx + tw - lw
                else:
                    lx = tx + marL
                asc = lh - fnt.size
                if bu and i == 0:
                    d.text((tx, cy + asc / 2), bu, font=fnt, fill=col)
                d.text((lx, cy + asc / 2), ln, font=fnt, fill=col)
                cy += lh
            cy += sa
        used = cy - start_y
        body = "".join(r.text for pp in tf.paragraphs for r in pp.runs)
        if body and used > (h - mt - mb) + 2 * SCALE:
            overflow.append((shp.shape_id, round((used) / SCALE), round(h / SCALE),
                             body[:34]))
    img.save(os.path.join(OUTDIR, "slide%02d.png" % idx))
    return overflow


prs = Presentation(SRC)
os.makedirs(OUTDIR, exist_ok=True)
bad = 0
for i, s in enumerate(prs.slides, 1):
    ov = render(s, i)
    for o in ov:
        bad += 1
        print("OVERFLOW slide%02d  used=%spx box=%spx  %r" % (i, o[1], o[2], o[3]))
print("rendered %d slides, %d overflow warnings" % (len(prs.slides._sldIdLst), bad))
