# -*- coding: utf-8 -*-
"""アニメーション付きDADS準拠PPTXの最小構成サンプル（表紙・順次表示・コード解説・出典の4枚）。

実行例:
    .venv/bin/python .claude/skills/dads-pptx/examples/build_anim_example.py /tmp/anim_example.pptx
    .venv/bin/python .claude/skills/dads-pptx/scripts/dads_anim.py /tmp/anim_example.pptx
"""
import sys, os

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))

from pptx import Presentation
from dads_theme import *   # noqa
from dads_anim import grab, Timeline, code_block

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "anim_example.pptx")
DECK_TITLE = "アニメーションのサンプル"
CONTENT_W = CANVAS_W - MARGIN_X * 2

prs = Presentation()
prs.slide_width, prs.slide_height = px(CANVAS_W), px(CANVAS_H)
BLANK = prs.slide_layouts[6]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, CANVAS_W, CANVAS_H, fill=bg)
    return s


def header(slide, kicker, title):
    rect(slide, MARGIN_X, 52, 6, 26, fill=PRIMARY)
    block(slide, MARGIN_X + 16, 52, CONTENT_W - 16, 26,
          [dict(text=kicker, style="Oln-14B-100", color=PRIMARY)], anchor="m")
    block(slide, MARGIN_X, 84, CONTENT_W, 48,
          [dict(text=title, style="Std-32B-150", color=INK)])
    rule(slide, MARGIN_X, 140, CONTENT_W, 2, BORDER)


def footer(slide, no):
    rule(slide, MARGIN_X, 660, CONTENT_W, 1, BORDER_SOFT)
    block(slide, MARGIN_X, 672, 700, 20,
          [dict(text=DECK_TITLE, style="Dns-14N-130", color=INK_MUTE)])
    block(slide, CANVAS_W - MARGIN_X - 200, 672, 200, 20,
          [dict(text=str(no), style="Dns-14B-130", color=INK_MUTE, align="r")])


# ---------------------------------------------------------------- 表紙（画面切り替えのみ）
s = new_slide(TERTIARY)
rect(s, 0, 0, 8, CANVAS_H, fill=SECONDARY)
block(s, MARGIN_X + 24, 246, 1000, 90,
      [dict(text="資料タイトル", style="Dsp-64B-140", color=WHITE)])
rect(s, MARGIN_X + 24, 372, 120, 4, fill=SECONDARY)
block(s, MARGIN_X + 24, 404, 860, 40,
      [dict(text="クリックで順に表示する資料", style="Std-24N-150", color=BLUE[100])])
Timeline(s).apply()

# ---------------------------------------------------------------- プロセスフロー（1クリック＝矢印＋カード）
s = new_slide()
header(s, "01 ── 順次表示", "手順をクリックごとに見せる")
tl = Timeline(s)
steps = [("手順1", "決める", "何を扱うかを決める。"),
         ("手順2", "作る", "決めた内容で形にする。"),
         ("手順3", "確かめる", "結果を見て直す。")]
gap = 48
cw = (CONTENT_W - gap * 2) / 3
for i, (tag, head, body) in enumerate(steps):
    x = MARGIN_X + i * (cw + gap)
    with grab(s) as g:
        if i:
            arrow_right(s, x - gap / 2, 330, 20, BORDER)
        rect(s, x, 200, cw, 260, fill=WHITE, line=BORDER, radius=RADIUS["none"])
        rect(s, x, 200, cw, 48, fill=PRIMARY)
        block(s, x + 24, 200, cw - 48, 48,
              [dict(text=tag, style="Std-20B-150", color=WHITE)], anchor="m")
        block(s, x + 24, 272, cw - 48, 160,
              [dict(text=head, style="Std-24B-150", color=INK),
               dict(text=body, style="Std-16N-170", color=INK_SUB, space_before=8)])
    tl.click(f"{tag}：{head}", g.shapes)
with grab(s) as g:
    rect(s, MARGIN_X, 496, CONTENT_W, 72, fill=BG_KEY)
    block(s, MARGIN_X + 32, 496, CONTENT_W - 64, 72,
          [dict(text="最後に結論の帯を出す。全表示の状態だけでも内容が伝わるようにする。",
                style="Std-20B-150", color=TERTIARY)], anchor="m")
tl.click("結論を表示", g.shapes)
footer(s, 2)
tl.apply()

# ---------------------------------------------------------------- コード解説（行ハイライト＋注釈）
s = new_slide()
header(s, "02 ── コード解説", "説明する行を順に強調する")
tl = Timeline(s)
cb = code_block(s, MARGIN_X, 200, 640, [
    "SELECT",
    "  store,",
    "  SUM(amount) AS total",
    "FROM sales",
    "GROUP BY store;",
], highlights=[(2, 2), (4, 4)])
notes = ["SUM で店舗ごとの合計を出す", "GROUP BY で店舗ごとに1行へまとめる"]
for i, text in enumerate(notes):
    with grab(s) as g:
        block(s, 752, 216 + i * 72, CONTENT_W - 688, 56,
              [dict(text=f"{i + 1}. {text}", style="Std-18N-160", color=INK)])
    tl.click(text, cb["highlights"][i] + g.shapes)
footer(s, 3)
tl.apply()

# ---------------------------------------------------------------- 出典
s = new_slide()
header(s, "APPENDIX", "本資料について（出典・注記）")
block(s, MARGIN_X, 200, CONTENT_W, 40,
      [dict(text="アニメーションはフェードのみ・クリック進行です。動きがなくても内容が伝わる構成にしています。",
            style="Std-16N-170", color=INK)])
rule(s, MARGIN_X, 470, CONTENT_W, 2, BORDER)
block(s, MARGIN_X, 492, CONTENT_W, 60,
      [dict(text="出典：デジタル庁デザインシステムウェブサイト https://design.digital.go.jp/dads/ のコンテンツを加工して作成",
            style="Std-16N-170", color=INK_SUB),
       dict(text="本資料はデジタル庁が作成・監修したものではありません。",
            style="Dns-14N-130", color=INK_MUTE)])
footer(s, 4)
Timeline(s).apply()

prs.save(OUT)
print("saved:", OUT)
