# -*- coding: utf-8 -*-
"""AI駆動開発 プレゼンテーション（DADS 基本デザイン準拠）を生成する。"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR
from dads_theme import *   # noqa

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "out", "AI駆動開発_DADS.pptx")

DECK_TITLE = "AI駆動開発"
CONTENT_W = CANVAS_W - MARGIN_X * 2          # 1152
BODY_TOP = 168                               # 見出しブロックの下端
BODY_BOTTOM = 648                            # フッター上端
BODY_H = BODY_BOTTOM - BODY_TOP

prs = Presentation()
prs.slide_width, prs.slide_height = px(CANVAS_W), px(CANVAS_H)
BLANK = prs.slide_layouts[6]


def new_slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, CANVAS_W, CANVAS_H, fill=bg)
    return s


def header(slide, kicker, title, lead=None):
    """カテゴリーラベル・見出し・リード文・ディバイダーからなる見出しブロック。"""
    rect(slide, MARGIN_X, 52, 6, 26, fill=PRIMARY)          # キーカラーのアクセント
    block(slide, MARGIN_X + 16, 52, CONTENT_W - 16, 26,
          [dict(text=kicker, style="Oln-14B-100", color=PRIMARY)], anchor="m")
    block(slide, MARGIN_X, 84, CONTENT_W, 48,
          [dict(text=title, style="Std-32B-150", color=INK)])
    rule(slide, MARGIN_X, 140, CONTENT_W, 2, BORDER)
    if lead:
        block(slide, MARGIN_X, 152, CONTENT_W, 30,
              [dict(text=lead, style="Std-17N-170", color=INK_SUB)])


def footer(slide, no):
    rule(slide, MARGIN_X, 660, CONTENT_W, 1, BORDER_SOFT)
    block(slide, MARGIN_X, 672, 700, 20,
          [dict(text=DECK_TITLE, style="Dns-14N-130", color=INK_MUTE)])
    block(slide, CANVAS_W - MARGIN_X - 200, 672, 200, 20,
          [dict(text=str(no), style="Dns-14B-130", color=INK_MUTE, align="r")])


def card(slide, x, y, w, h, fill=SURFACE, line=BORDER, radius=RADIUS["m"]):
    return rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)


def datatable(slide, x, y, w, cols, rows, row_h=54, head_h=48):
    """DADSのテーブルに倣った横罫線ベースの表。"""
    widths = [w * c for c in cols["ratio"]]
    xs, acc = [], x
    for cw in widths:
        xs.append(acc)
        acc += cw
    rect(slide, x, y, w, head_h, fill=BG_KEY)
    for i, label in enumerate(cols["labels"]):
        block(slide, xs[i] + 16, y, widths[i] - 32, head_h,
              [dict(text=label, style="Dns-14B-130", color=PRIMARY)], anchor="m")
    rule(slide, x, y + head_h, w, 2, PRIMARY)
    cy = y + head_h + 2
    for r in rows:
        for i, cell in enumerate(r):
            style = "Dns-16N-130" if i else "Dns-16N-130"
            color = INK
            bold_style = "Dns-14B-130"
            block(slide, xs[i] + 16, cy, widths[i] - 32, row_h,
                  [dict(text=cell, style=("Std-16B-170" if i == 0 else "Std-16N-170"),
                        color=(INK if i == 0 else INK_SUB))], anchor="m")
        rule(slide, x, cy + row_h, w, 1, BORDER_SOFT)
        cy += row_h + 1
    return cy


# ==============================================================  01 表紙
s = new_slide(TERTIARY)
rect(s, 0, 0, 8, CANVAS_H, fill=SECONDARY)
block(s, MARGIN_X + 24, 200, 900, 40,
      [dict(text="AI-DRIVEN DEVELOPMENT", style="Oln-16B-100", color=BLUE[300])])
block(s, MARGIN_X + 24, 246, 1000, 90,
      [dict(text="AI駆動開発", style="Dsp-64B-140", color=WHITE)])
rect(s, MARGIN_X + 24, 372, 120, 4, fill=SECONDARY)
block(s, MARGIN_X + 24, 404, 860, 80,
      [dict(text="開発プロセス全体にAIを組み込むための", style="Std-24N-150", color=BLUE[100]),
       dict(text="考え方・進め方・リスク対策", style="Std-24N-150", color=BLUE[100])])
rule(s, MARGIN_X + 24, 600, CONTENT_W - 48, 1, BLUE[1100])
block(s, MARGIN_X + 24, 620, 600, 24,
      [dict(text="社内勉強会資料", style="Dns-16N-130", color=BLUE[200])])
block(s, CANVAS_W - MARGIN_X - 424, 620, 400, 24,
      [dict(text="2026年9月", style="Dns-16N-130", color=BLUE[200], align="r")])

# ==============================================================  02 目次
s = new_slide()
header(s, "AGENDA", "目次")
items = [
    ("01", "AI駆動開発とは", "定義と、従来の開発との違い"),
    ("02", "取り組む背景", "いま検討する理由と、前提の整理"),
    ("03", "適用領域とツール", "工程ごとの使いどころと選定観点"),
    ("04", "導入の進め方", "4フェーズのロードマップとKPI"),
    ("05", "品質・リスク・ガバナンス", "検証の仕組みとルール整備"),
    ("06", "組織と役割の変化", "スキルシフトとアンチパターン"),
]
gy, gh, gap = 188, 112, 24
for i, (no, t1, t2) in enumerate(items):
    col, row = i // 3, i % 3
    x = MARGIN_X + col * (CONTENT_W / 2 + 16)
    y = gy + row * (gh + gap)
    w = CONTENT_W / 2 - 16
    card(s, x, y, w, gh)
    rect(s, x, y + 20, 4, gh - 40, fill=PRIMARY)
    block(s, x + 28, y, 56, gh,
          [dict(text=no, style="Std-28B-150", color=PRIMARY)], anchor="m")
    block(s, x + 92, y, w - 116, gh, anchor="m", lines=
          [dict(text=t1, style="Std-20B-150", color=INK),
           dict(text=t2, style="Std-16N-170", color=INK_MUTE, space_before=4)])
footer(s, 2)

# ==============================================================  03 定義
s = new_slide()
header(s, "01 ── AI駆動開発とは", "AIを前提に開発プロセスを組み直すこと")
card(s, MARGIN_X, 192, CONTENT_W, 96, fill=BG_KEY, line=None, radius=RADIUS["m"])
block(s, MARGIN_X + 32, 192, CONTENT_W - 64, 96,
      [dict(text="生成AI・AIエージェントを開発の主要な作業単位として組み込み、要件定義から運用まで",
            style="Std-20B-150", color=TERTIARY),
       dict(text="の各工程を「AIが下書きし、人が判断・検証する」形へ再設計する開発スタイル。",
            style="Std-20B-150", color=TERTIARY)], anchor="m")
cards = [
    ("コード生成だけではない", "要件整理、設計案の比較、テスト設計、レビュー、ドキュメント、運用時の調査まで対象になる。"),
    ("人の役割は「決める・確かめる」", "書く時間が減る分を、課題定義・受け入れ基準の設計・検証に再配分する。"),
    ("前提は再現可能なプロセス", "仕様・テスト・CIが整っているほど、AIの出力品質と安全性が安定する。"),
]
cw = (CONTENT_W - 32) / 3
for i, (t1, t2) in enumerate(cards):
    x = MARGIN_X + i * (cw + 16)
    card(s, x, 312, cw, 186, radius=RADIUS["none"])
    rect(s, x, 312, cw, 6, fill=PRIMARY, radius=0)
    block(s, x + 24, 344, cw - 48, 140,
          [dict(text=t1, style="Std-20B-150", color=INK),
           dict(text=t2, style="Std-16N-170", color=INK_SUB, space_before=8)])
block(s, MARGIN_X, 524, CONTENT_W, 60,
      [dict(text="ツールを配布することがAI駆動開発ではない。作業の分担と検証の方法を定義して、はじめて成立する。",
            style="Std-17N-170", color=INK_SUB, bullet="※")])
footer(s, 3)

# ==============================================================  04 従来との違い
s = new_slide()
header(s, "01 ── AI駆動開発とは", "従来の開発との違い",
       "ボトルネックが「実装速度」から「仕様の曖昧さとレビュー容量」へ移る。")
datatable(s, MARGIN_X, 200, CONTENT_W,
          {"labels": ["観点", "従来の開発", "AI駆動開発"], "ratio": [0.22, 0.39, 0.39]},
          [["成果物の起点", "人が一次ドラフトを作成する", "AIが一次ドラフト、人が確定させる"],
           ["主なボトルネック", "実装そのものの速度", "仕様の曖昧さ・レビューの処理量"],
           ["品質の担保", "レビューとテスト", "自動テストと静的解析＋人のレビュー"],
           ["求められるスキル", "実装力・設計力", "課題定義力・検証設計力・文脈の言語化"],
           ["コスト構造", "人月が中心", "人月＋AI利用料（従量）"]],
          row_h=58)
card(s, MARGIN_X, 556, CONTENT_W, 68, fill=BG_KEY, line=None)
block(s, MARGIN_X + 32, 556, CONTENT_W - 64, 68,
      [dict(text="違いの本質は「速く書けること」ではなく、判断と検証にかける時間の配分が変わること。",
            style="Std-17N-170", color=TERTIARY)], anchor="m")
footer(s, 4)

# ==============================================================  05 背景
s = new_slide()
header(s, "02 ── 取り組む背景", "いま検討する理由")
bgs = [
    ("01", "モデルと実行環境の進歩",
     ["長い文脈を扱えるようになり、リポジトリ全体を前提とした作業が可能になった。",
      "計画・実行・検証を繰り返すエージェント型の利用が現実的になった。"]),
    ("02", "開発ワークフローへの統合",
     ["IDE・CI・課題管理に組み込まれ、日常の作業導線の中で使えるようになった。",
      "利用ログや権限管理など、組織で運用するための機能が整ってきた。"]),
    ("03", "増員に頼らない生産性の要求",
     ["人材の確保が難しく、既存チームの生産性向上が現実的な選択肢になっている。",
      "レガシー資産の保守・移行など、労力の大きい領域が残っている。"]),
]
y = 190
for no, t1, lines in bgs:
    card(s, MARGIN_X, y, CONTENT_W, 116)
    circle(s, MARGIN_X + 28, y + 34, 48, PRIMARY)
    block(s, MARGIN_X + 28, y + 34, 48, 48,
          [dict(text=no, style="Std-20B-150", color=WHITE, align="c")], anchor="m")
    block(s, MARGIN_X + 100, y + 18, CONTENT_W - 132, 92,
          [dict(text=t1, style="Std-20B-150", color=INK)] +
          [dict(text=l, style="Std-16N-170", color=INK_SUB, bullet="・") for l in lines])
    y += 132
card(s, MARGIN_X, 592, CONTENT_W, 56, fill=BG_KEY, line=None, radius=RADIUS["s"])
block(s, MARGIN_X + 24, 592, CONTENT_W - 48, 56,
      [dict(text="注意：「導入すれば速くなる」わけではない。効果は対象工程の選び方と運用ルールの設計に依存する。",
            style="Std-17N-170", color=TERTIARY)], anchor="m")
footer(s, 5)

# ==============================================================  06 適用領域
s = new_slide()
header(s, "03 ── 適用領域とツール", "開発プロセスにおける適用領域",
       "上流ほど「案出しと整理」、下流ほど「生成と検査」に効く。")
steps = [
    ("要件定義", "論点の洗い出し／ユーザーストーリーの草案／曖昧な記述の検出"),
    ("設計", "方式の比較検討／インターフェース定義／設計判断記録のたたき台"),
    ("実装", "コード生成／リファクタリング／既存資産の移行・移植"),
    ("テスト", "テスト観点の列挙／単体テストの生成／テストデータの作成"),
    ("レビュー", "差分の要約／観点別の指摘／静的解析結果の整理"),
    ("運用・保守", "障害調査の補助／ログ要約／ドキュメントの追随更新"),
]
cw = (CONTENT_W - 5 * 12) / 6
for i, (t1, t2) in enumerate(steps):
    x = MARGIN_X + i * (cw + 12)
    card(s, x, 224, cw, 200, radius=RADIUS["none"])
    rect(s, x, 224, cw, 40, fill=PRIMARY, radius=0)
    block(s, x, 224, cw, 40,
          [dict(text=t1, style="Oln-16B-100", color=WHITE, align="c")], anchor="m")
    block(s, x + 16, 280, cw - 32, 160,
          [dict(text=t2, style="Dns-16N-130", color=INK_SUB)])
    if i < 5:
        arrow_right(s, x + cw + 6, 244, 10, BORDER)
card(s, MARGIN_X, 460, CONTENT_W, 140, fill=BG_KEY, line=None)
block(s, MARGIN_X + 32, 484, CONTENT_W - 64, 100,
      [dict(text="最初に着手する領域の選び方", style="Std-18B-160", color=TERTIARY),
       dict(text="影響範囲が閉じている／正解の判定が自動化できる／繰り返し発生する、の3条件を満たす作業から始める。",
            style="Std-16N-170", color=INK, space_before=4),
       dict(text="例：単体テストの追加、定型的な移行作業、レビュー前の差分要約、ドキュメントの更新。",
            style="Std-16N-170", color=INK)])
footer(s, 6)

# ==============================================================  07 ツール分類
s = new_slide()
header(s, "03 ── 適用領域とツール", "ツールの分類と選定の観点")
tools = [
    ("補完型", "インラインでの入力補完", "導入が容易で既存フローを変えない。効果は実装工程に限定される。"),
    ("対話型", "チャット／IDE統合", "設計相談・調査・影響範囲の把握に向く。文脈の与え方で品質が変わる。"),
    ("エージェント型", "自律的に複数手順を実行", "横断的な変更やテスト実行まで担う。権限と実行範囲の設計が要になる。"),
    ("自動化組込型", "CI／レビューへの組込", "レビュー補助やテスト生成を常時実行。属人性を減らせる。"),
]
cw = (CONTENT_W - 3 * 16) / 4
for i, (t1, t2, t3) in enumerate(tools):
    x = MARGIN_X + i * (cw + 16)
    card(s, x, 192, cw, 216)
    badge(s, x + 20, 216, t1, fill=BG_KEY, fg=PRIMARY, w=cw - 40)
    block(s, x + 20, 260, cw - 40, 130,
          [dict(text=t2, style="Std-18B-160", color=INK),
           dict(text=t3, style="Std-16N-170", color=INK_SUB, space_before=8)])
block(s, MARGIN_X, 436, CONTENT_W, 30,
      [dict(text="選定時に確認する4点", style="Std-20B-150", color=INK)])
checks = [
    ("入力データの扱い", "学習利用の有無、保存期間、リージョン"),
    ("権限とログ", "アカウント管理、操作ログ、実行範囲の制限"),
    ("既存環境との接続", "IDE・CI・課題管理・認証基盤との統合可否"),
    ("コストの可視性", "従量課金の単位、上限設定、部門別の把握"),
]
cw2 = (CONTENT_W - 3 * 16) / 4
for i, (t1, t2) in enumerate(checks):
    x = MARGIN_X + i * (cw2 + 16)
    card(s, x, 476, cw2, 132, fill=WHITE, line=BORDER)
    rect(s, x + 20, 500, 12, 12, fill=PRIMARY, radius=2)
    block(s, x + 20, 522, cw2 - 40, 76,
          [dict(text=t1, style="Std-17N-170", color=INK),
           dict(text=t2, style="Dns-14N-130", color=INK_MUTE, space_before=4)])
footer(s, 7)

# ==============================================================  08 導入ステップ
s = new_slide()
header(s, "04 ── 導入の進め方", "4フェーズのロードマップ",
       "小さく試し、記録を残し、ルールとセットで広げる。")
phases = [
    ("Phase 0", "準備", "〜1か月", GRAY[700],
     ["利用ルールと対象範囲の決定", "利用するツールと権限の整備", "現状のベースライン計測"]),
    ("Phase 1", "試行", "1〜3か月", "0031D8",
     ["少人数・低リスク領域で検証", "うまくいった使い方を記録", "阻害要因の洗い出し"]),
    ("Phase 2", "定着", "3〜6か月", PRIMARY,
     ["定型プロンプトと事例の共有", "CI・レビューへの組込", "教育とオンボーディング"]),
    ("Phase 3", "拡大", "6か月〜", TERTIARY,
     ["適用工程を上流・運用へ拡大", "KPIによる継続的な見直し", "ルールと体制の更新"]),
]
cw = (CONTENT_W - 3 * 20) / 4
for i, (ph, name, term, col, lines) in enumerate(phases):
    x = MARGIN_X + i * (cw + 20)
    card(s, x, 208, cw, 292, radius=RADIUS["none"])
    rect(s, x, 208, cw, 76, fill=col, radius=0)
    block(s, x + 24, 208, cw - 48, 76,
          [dict(text=ph, style="Oln-14B-100", color=WHITE),
           dict(text=name, style="Std-24B-150", color=WHITE, space_before=4)], anchor="m")
    block(s, x + 24, 300, cw - 48, 24,
          [dict(text=term, style="Dns-14B-130", color=INK_MUTE)])
    block(s, x + 24, 332, cw - 48, 156,
          [dict(text=l, style="Std-16N-170", color=INK_SUB, bullet="・") for l in lines])
    if i < 3:
        arrow_right(s, x + cw + 10, 246, 12, BORDER)
card(s, MARGIN_X, 524, CONTENT_W, 100, fill=BG_KEY, line=None)
block(s, MARGIN_X + 32, 544, CONTENT_W - 64, 76,
      [dict(text="Phase 0 で必ずやること：ベースラインの計測", style="Std-18B-160", color=TERTIARY),
       dict(text="導入前のリードタイム・変更失敗率・レビュー待ち時間を記録しておかないと、後から効果を説明できない。",
            style="Std-16N-170", color=INK, space_before=4)])
footer(s, 8)

# ==============================================================  09 KPI
s = new_slide()
header(s, "04 ── 導入の進め方", "効果測定：KPIは速度と品質を対で置く")
datatable(s, MARGIN_X, 196, CONTENT_W,
          {"labels": ["観点", "指標の例", "運用上の注意"], "ratio": [0.18, 0.40, 0.42]},
          [["スピード", "変更のリードタイム／レビュー待ち時間", "工程別に分けて見ないと改善点が特定できない"],
           ["品質", "変更失敗率／障害の復旧時間", "速度指標と必ず同時に確認する"],
           ["量", "変更件数／自動テストの網羅範囲", "件数だけを目標にすると分割の乱用を招く"],
           ["開発者体験", "満足度調査／手戻りの発生率", "定量指標を定性調査で補う"],
           ["コスト", "変更1件あたりのAI利用料", "効果と対で見て、上限を設定する"]],
          row_h=56)
card(s, MARGIN_X, 552, CONTENT_W, 60, fill=WHITE, line=WARN_ORANGE, radius=RADIUS["none"])
rect(s, MARGIN_X, 552, 6, 60, fill=WARN_ORANGE)
block(s, MARGIN_X + 28, 552, CONTENT_W - 56, 60,
      [dict(text="速度指標のみを追うと、品質の劣化が数か月遅れて表面化する。導入初期ほど品質指標を主に見る。",
            style="Std-17N-170", color=INK)], anchor="m")
footer(s, 9)

# ==============================================================  10 品質担保
s = new_slide()
header(s, "05 ── 品質・リスク・ガバナンス", "品質を担保する3つの層")
layers = [
    ("入力の管理", "AIに渡す前提を揃える",
     ["仕様・制約・既存の設計方針を明示する", "参照させる範囲を限定する", "禁止事項を先に伝える"]),
    ("出力の検証", "人とツールの二重で確かめる",
     ["自動テストと静的解析を必ず通す", "生成コードも通常と同じレビュー基準", "根拠と代替案の提示を求める"]),
    ("記録の保全", "後から説明できる状態にする",
     ["生成の経緯とレビュー結果を残す", "利用ログを保管する", "判断の理由を記録に残す"]),
]
cw = (CONTENT_W - 2 * 20) / 3
for i, (t1, t2, lines) in enumerate(layers):
    x = MARGIN_X + i * (cw + 20)
    card(s, x, 200, cw, 260)
    block(s, x + 24, 224, cw - 48, 30,
          [dict(text="LAYER %d" % (i + 1), style="Oln-14B-100", color=PRIMARY)])
    block(s, x + 24, 256, cw - 48, 60,
          [dict(text=t1, style="Std-24B-150", color=INK),
           dict(text=t2, style="Dns-14N-130", color=INK_MUTE)])
    rule(s, x + 24, 332, cw - 48, 1, BORDER_SOFT)
    block(s, x + 24, 348, cw - 48, 100,
          [dict(text=l, style="Std-16N-170", color=INK_SUB, bullet="・") for l in lines])
card(s, MARGIN_X, 484, CONTENT_W, 128, fill=WHITE, line=BORDER)
block(s, MARGIN_X + 32, 508, CONTENT_W - 64, 100,
      [dict(text="レビューが律速になることを前提に設計する", style="Std-20B-150", color=INK),
       dict(text="生成量が増えるとレビュー待ちが最大の停滞要因になる。差分を小さく保つ、観点を自動化する、",
            style="Std-16N-170", color=INK_SUB, space_before=4),
       dict(text="レビュー担当を分散させる、の3点を導入と同時に手当てする。",
            style="Std-16N-170", color=INK_SUB)])
footer(s, 10)

# ==============================================================  11 リスク
s = new_slide()
header(s, "05 ── 品質・リスク・ガバナンス", "主なリスクと対策")
risks = [
    ("情報漏えい", ERROR,
     ["入力してよい情報の分類を定める", "送信先・保存期間・学習利用の有無を確認", "業務データは管理された基盤に限定"]),
    ("誤りの混入（ハルシネーション）", WARN_ORANGE,
     ["出力は必ず実行・テストで検証", "存在しないAPIや依存の確認を自動化", "確信度ではなく根拠で判断する"]),
    ("ライセンス・権利", WARN_YELLOW,
     ["依存関係とライセンスを自動スキャン", "生成物の扱いを規程に明記", "外部提供物は特に確認を厚く"]),
    ("スキルの空洞化", SUCCESS,
     ["意図を説明できることをレビュー要件に", "基礎の教育計画とセットで進める", "重要領域は人の設計を先行させる"]),
]
cw = (CONTENT_W - 20) / 2
for i, (t1, col, lines) in enumerate(risks):
    x = MARGIN_X + (i % 2) * (cw + 20)
    y = 196 + (i // 2) * 214
    card(s, x, y, cw, 190, radius=RADIUS["none"])
    rect(s, x, y, 6, 190, fill=col)
    block(s, x + 32, y + 24, cw - 64, 32,
          [dict(text=t1, style="Std-20B-150", color=INK)])
    rule(s, x + 32, y + 66, cw - 64, 1, BORDER_SOFT)
    block(s, x + 32, y + 82, cw - 64, 96,
          [dict(text=l, style="Std-16N-170", color=INK_SUB, bullet="・") for l in lines])
footer(s, 11)

# ==============================================================  12 ガバナンス
s = new_slide()
header(s, "05 ── 品質・リスク・ガバナンス", "ルールとガバナンスの整備",
       "現場が守れる粒度で定め、四半期ごとに見直す。")
left = ("あらかじめ定めること",
        ["利用を許可するツールとアカウントの管理方法",
         "入力してよい情報／禁止する情報の分類",
         "生成物の権利・責任の所在と表示の要否",
         "ログの取得範囲と保管期間",
         "ルールから外れる場合の申請と承認の経路"])
right = ("運用として回すこと",
         ["相談窓口を1か所に集約し、判断のばらつきを防ぐ",
          "四半期ごとにルールと対象ツールを見直す",
          "インシデント発生時の報告経路を明示する",
          "利用状況とコストを部門単位で可視化する",
          "良い使い方の事例を継続的に共有する"])
cw = (CONTENT_W - 24) / 2
for i, (t1, lines) in enumerate([left, right]):
    x = MARGIN_X + i * (cw + 24)
    card(s, x, 200, cw, 320, radius=RADIUS["none"])
    rect(s, x, 200, cw, 56, fill=(PRIMARY if i == 0 else GRAY[700]), radius=0)
    block(s, x + 24, 200, cw - 48, 56,
          [dict(text=t1, style="Std-20B-150", color=WHITE)], anchor="m")
    for j, l in enumerate(lines):
        yy = 276 + j * 46
        rect(s, x + 24, yy + 8, 10, 10, fill=(PRIMARY if i == 0 else GRAY[700]), radius=2)
        block(s, x + 46, yy, cw - 70, 40,
              [dict(text=l, style="Std-16N-170", color=INK)])
card(s, MARGIN_X, 544, CONTENT_W, 68, fill=BG_KEY, line=None)
block(s, MARGIN_X + 32, 544, CONTENT_W - 64, 68,
      [dict(text="禁止事項を増やしすぎると、記録に残らない利用（シャドー利用）を招き、かえって統制が効かなくなる。",
            style="Std-17N-170", color=TERTIARY)], anchor="m")
footer(s, 12)

# ==============================================================  13 役割の変化
s = new_slide()
header(s, "06 ── 組織と役割の変化", "求められる力の移り変わり")
cols = [("これまで", GRAY[700], BORDER,
         ["実装した量が成果として見えやすい",
          "レビューは最終工程での確認",
          "設計の意図は個人の頭の中に残る",
          "見積りは作業量から積み上げる"]),
        ("これから", PRIMARY, PRIMARY,
         ["何を作らないかの判断が成果になる",
          "レビューが律速。観点の自動化と分散が要る",
          "文脈を明文化して初めてAIが使える資産になる",
          "見積りは検証コストを含めて考える"])]
cw = (CONTENT_W - 56) / 2
for i, (t1, col, bcol) in enumerate([(c[0], c[1], c[2]) for c in cols]):
    x = MARGIN_X + i * (cw + 56)
    lines = cols[i][3]
    card(s, x, 200, cw, 268, line=bcol)
    block(s, x + 28, 224, cw - 56, 40,
          [dict(text=t1, style="Std-24B-150", color=col)])
    rule(s, x + 28, 276, cw - 56, 2, col)
    block(s, x + 28, 296, cw - 56, 150,
          [dict(text=l, style="Std-17N-170", color=INK, bullet="・", space_after=8)
           for l in lines])
arrow_right(s, MARGIN_X + cw + 28, 330, 20, BORDER)
block(s, MARGIN_X, 492, CONTENT_W, 30,
      [dict(text="陥りやすい進め方", style="Std-20B-150", color=INK)])
anti = [("目的のないツール配布", "配って終わりでは効果が測れず、使う人と使わない人に分かれて定着しない。"),
        ("速度だけのKPI", "短期の指標は改善して見えるが、品質の劣化が遅れて表面化する。"),
        ("例外の多すぎるルール", "実態に合わないルールは守られず、記録に残らない利用を増やす。")]
cw3 = (CONTENT_W - 32) / 3
for i, (t1, t2) in enumerate(anti):
    x = MARGIN_X + i * (cw3 + 16)
    card(s, x, 530, cw3, 108, fill=WHITE, line=BORDER, radius=RADIUS["s"])
    block(s, x + 20, 550, cw3 - 40, 76,
          [dict(text="×　" + t1, style="Std-17N-170", color=ERROR),
           dict(text=t2, style="Dns-14N-130", color=INK_SUB, space_before=4)])
footer(s, 13)

# ==============================================================  14 まとめ
s = new_slide()
header(s, "SUMMARY", "まとめ")
msgs = [("01", "ツール導入ではなく、工程の再設計として扱う",
         "誰が下書きし、誰が判断し、どこで検証するかを工程ごとに決める。"),
        ("02", "効果は仕様の明確さと検証の仕組みに比例する",
         "曖昧な仕様と手動頼りの検証のままでは、生成量が増えても成果にならない。"),
        ("03", "速度と品質のKPIを対に置き、小さく試して広げる",
         "低リスク領域で計測しながら、ルールと教育をセットで展開する。")]
y = 190
for no, t1, t2 in msgs:
    card(s, MARGIN_X, y, CONTENT_W, 110, fill=BG_KEY, line=None, radius=RADIUS["none"])
    rect(s, MARGIN_X, y, 6, 110, fill=PRIMARY)
    block(s, MARGIN_X + 32, y, 60, 110,
          [dict(text=no, style="Std-28B-150", color=PRIMARY)], anchor="m")
    block(s, MARGIN_X + 100, y + 24, CONTENT_W - 132, 70,
          [dict(text=t1, style="Std-22B-150", color=TERTIARY),
           dict(text=t2, style="Std-16N-170", color=INK, space_before=4)])
    y += 126
block(s, MARGIN_X, 562, CONTENT_W, 30,
      [dict(text="次のアクション", style="Std-20B-150", color=INK)])
acts = ["対象領域を1つ選ぶ", "利用ルールの草案を作る", "ベースラインを計測する"]
cw = (CONTENT_W - 32) / 3
for i, a in enumerate(acts):
    x = MARGIN_X + i * (cw + 16)
    card(s, x, 596, cw, 48, fill=WHITE, line=PRIMARY, radius=RADIUS["s"])
    block(s, x, 596, cw, 48,
          [dict(text=a, style="Std-17N-170", color=PRIMARY, align="c")], anchor="m")
footer(s, 14)

# ==============================================================  15 出典
s = new_slide()
header(s, "APPENDIX", "本資料について（出典・注記）")
notes = [
    ("デザイン", ["本資料のレイアウト・配色・書体・余白は、デジタル庁デザインシステムβ版 v2.17.1 の",
                "基本デザイン（カラー、タイポグラフィ、余白、角の形状、エレベーション）を参考に構成しています。"]),
    ("カラー", ["プリミティブカラー Blue（キーカラー）とニュートラルカラー Solid Gray を使用。",
              "テキストは4.5:1以上、罫線などの非テキスト要素は3:1以上のコントラスト比を確保しています。"]),
    ("書体", ["Noto Sans JP（SIL Open Font License 1.1）。テキストスタイルはDADSのトークン定義に準拠。",
             "閲覧環境に当該書体がない場合は代替書体で表示されます。"]),
    ("内容", ["AI駆動開発に関する一般的な整理であり、特定の製品・サービスの推奨ではありません。",
             "実際の導入は、各組織の情報管理規程および契約条件に従ってください。"]),
]
y = 196
for t1, lines in notes:
    badge(s, MARGIN_X, y + 6, t1, fill=BG_KEY, fg=PRIMARY, w=104)
    block(s, MARGIN_X + 128, y, CONTENT_W - 128, 70,
          [dict(text=l, style="Std-16N-170", color=INK) for l in lines])
    y += 94
rule(s, MARGIN_X, 578, CONTENT_W, 2, BORDER)
block(s, MARGIN_X, 596, CONTENT_W, 60,
      [dict(text="出典：デジタル庁デザインシステムウェブサイト https://design.digital.go.jp/dads/ のコンテンツを加工して作成",
            style="Std-16N-170", color=INK_SUB),
       dict(text="本資料はデジタル庁が作成・監修したものではありません。",
            style="Dns-14N-130", color=INK_MUTE)])
footer(s, 15)

prs.save(OUT)
print("saved:", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
