# -*- coding: utf-8 -*-
"""デジタル庁デザインシステム（DADS）β版の基本デザインをPPTXへ写像するための共通モジュール。

- カラー: @digital-go-jp/design-tokens v2.0.1 のプリミティブ/ニュートラルカラー
- タイポグラフィ: Noto Sans JP / テキストスタイル（Dsp, Std, Dns, Oln）
- 余白: 基準単位 8 CSS px の倍率スケール
- 角の形状: 角丸なし(0) / スモール(8) / ミディアム(12,16) / ラージ(24)

スライドは 1280 x 720 CSS px のキャンバスとして扱い、1 CSS px = 1/96 inch で配置する。
このとき font-size の CSS px は pt に 0.75 倍で対応する（16 px = 12 pt）。
"""
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- 単位変換
EMU_PER_PX = 914400 / 96          # CSS px -> EMU（96 CSS px = 1 inch）
CANVAS_W, CANVAS_H = 1280, 720    # 16:9


def px(v):
    """CSS px を EMU に変換する。"""
    return Emu(int(round(v * EMU_PER_PX)))


def fpt(css_px):
    """font-size の CSS px を pt に変換する。"""
    return Pt(css_px * 0.75)


# ---------------------------------------------------------------- カラー
# プリミティブカラー（青）
BLUE = {
    50: "E8F1FE", 100: "D9E6FF", 200: "C5D7FB", 300: "9DB7F9", 400: "7096F8",
    500: "4979F5", 600: "3460FB", 700: "264AF4", 800: "0031D8", 900: "0017C1",
    1000: "00118F", 1100: "000071", 1200: "000060",
}
# ニュートラルカラー（Solid Gray）
GRAY = {
    50: "F2F2F2", 100: "E6E6E6", 200: "CCCCCC", 300: "B3B3B3", 400: "999999",
    420: "949494", 500: "7F7F7F", 536: "767676", 600: "666666", 700: "4D4D4D",
    800: "333333", 900: "1A1A1A",
}
WHITE, BLACK = "FFFFFF", "000000"

# キーカラー（青のキーカラー例に相当する割り当て）
PRIMARY = BLUE[900]        # プライマリー: 主要動線・強調
SECONDARY = BLUE[500]      # セカンダリー: 副次的な強調（非テキスト用途）
TERTIARY = BLUE[1000]      # ターシャリー: 濃い面
BG_KEY = BLUE[50]          # バックグラウンドカラー

# 共通カラー（テキスト・境界）
INK = GRAY[900]            # 本文テキスト
INK_SUB = GRAY[700]        # 補足テキスト
INK_MUTE = GRAY[536]       # 4.5:1 を満たす最も明るいグレー
BORDER = GRAY[420]         # 白に対して3:1を満たす最も明るいグレー（非テキスト要件）
BORDER_SOFT = GRAY[300]    # 罫線内の補助線など、意味を持たない装飾用
SURFACE = WHITE

# セマンティックカラー
SUCCESS = "197A4B"         # green-800
ERROR = "EC0000"           # red-800
WARN_ORANGE = "C74700"     # orange-800
WARN_YELLOW = "B78F00"     # yellow-700

# ---------------------------------------------------------------- 余白 / 角丸
UNIT = 8
SP = {k: UNIT * k for k in range(1, 13)}   # 8 の倍率スケール
MARGIN_X = 64              # 基準単位の 8 倍
RADIUS = {"none": 0, "s": 8, "m": 12, "m_sq": 16, "l": 24}

# ---------------------------------------------------------------- 書体
FONT = "Noto Sans JP"
FONT_MONO = "Noto Sans Mono"

# テキストスタイル（DADSのタイポグラフィトークン）: (font-size px, bold, line-height)
STYLE = {
    "Dsp-64B-140": (64, True, 1.40),
    "Dsp-48B-140": (48, True, 1.40),
    "Std-36B-140": (36, True, 1.40),
    "Std-32B-150": (32, True, 1.50),
    "Std-28B-150": (28, True, 1.50),
    "Std-24B-150": (24, True, 1.50),
    "Std-24N-150": (24, False, 1.50),
    "Std-22B-150": (22, True, 1.50),
    "Std-20B-150": (20, True, 1.50),
    "Std-20N-150": (20, False, 1.50),
    "Std-18B-160": (18, True, 1.60),
    "Std-18N-160": (18, False, 1.60),
    "Std-17N-170": (17, False, 1.70),
    "Std-16B-170": (16, True, 1.70),
    "Std-16N-170": (16, False, 1.70),
    "Dns-16N-130": (16, False, 1.30),
    "Dns-14B-130": (14, True, 1.30),
    "Dns-14N-130": (14, False, 1.30),
    "Oln-16B-100": (16, True, 1.00),
    "Oln-14B-100": (14, True, 1.00),
    "Oln-14N-100": (14, False, 1.00),
}


# ---------------------------------------------------------------- コントラスト検証
def _lin(c):
    c = c / 255.0
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def contrast(hex_a, hex_b):
    """2色のコントラスト比を返す（WCAG 2.x の相対輝度による）。"""
    def lum(h):
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
        return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)
    la, lb = lum(hex_a), lum(hex_b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


# ---------------------------------------------------------------- 描画ヘルパ
def _apply_font(run, size_px, bold, color, font=FONT):
    """run に書体を適用する。CSS px 指定・和文書体（ea/cs）まで設定する。"""
    f = run.font
    f.size = fpt(size_px)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = font                      # latin
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):       # 和文・複合文字用の書体も揃える
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def set_bullet(p, marker="・", indent_px=20):
    """段落に行頭記号とぶら下げインデントを設定する。"""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent_px * EMU_PER_PX)))
    pPr.set("indent", str(-int(indent_px * EMU_PER_PX)))
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": FONT})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": marker})
    pPr.append(buFont)
    pPr.append(buChar)


def textbox(slide, x, y, w, h, anchor="t"):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    return tb


def para(tf, text, style="Std-18N-160", color=INK, align="l",
         space_before=0, space_after=0, first=False, bullet=None, font=FONT):
    """テキストフレームに段落を追加する。style は DADS のテキストスタイル名。"""
    size_px, bold, lh = STYLE[style]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                   "r": PP_ALIGN.RIGHT}[align]
    p.line_spacing = lh
    if space_before:
        p.space_before = fpt(space_before)
    if space_after:
        p.space_after = fpt(space_after)
    if bullet:
        set_bullet(p, bullet)
    run = p.add_run()
    run.text = text
    _apply_font(run, size_px, bold, color, font=font)
    return p


def block(slide, x, y, w, h, lines, anchor="t"):
    """(text, style, color, ...) のリストからテキストブロックを作る。"""
    tb = textbox(slide, x, y, w, h, anchor=anchor)
    tf = tb.text_frame
    for i, ln in enumerate(lines):
        para(tf, first=(i == 0), **ln)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_px=1, radius=0):
    """矩形／角丸矩形を配置する。radius は CSS px 相当の角丸半径。"""
    if radius:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     px(x), px(y), px(w), px(h))
        shp.adjustments[0] = radius / float(min(w, h))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     px(x), px(y), px(w), px(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(line_px * 0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False          # エレベーションは既定で高さレベル0
    shp.text_frame.word_wrap = True
    return shp


def circle(slide, x, y, d, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(d), px(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def arrow_right(slide, cx, cy, size, color=BORDER):
    """右向きの三角形。書体に依存しない矢印記号として使う。"""
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 px(cx - size / 2), px(cy - size / 2),
                                 px(size), px(size))
    shp.rotation = 90
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def rule(slide, x, y, w, thickness=1, color=BORDER):
    """ディバイダー（罫線）。"""
    return rect(slide, x, y, w, thickness, fill=color)


def badge(slide, x, y, text, fill=PRIMARY, fg=WHITE, style="Oln-14B-100",
          pad_x=12, h=28, radius=RADIUS["s"], w=None):
    """チップ／ラベル。"""
    size_px, _, _ = STYLE[style]
    est_w = w if w else pad_x * 2 + len(text) * size_px * 0.62
    shp = rect(slide, x, y, est_w, h, fill=fill, radius=radius)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, style=style, color=fg, align="c", first=True)
    return shp
