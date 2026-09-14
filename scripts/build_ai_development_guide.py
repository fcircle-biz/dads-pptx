# -*- coding: utf-8 -*-
"""AI駆動開発ガイドを生成する。既存のDADSヘルパを利用する。

実行: .venv/bin/python scripts/build_ai_development_guide.py
検証: .venv/bin/python scripts/render_check.py out/AI駆動開発ガイド/AI駆動開発ガイド_DADS.pptx out/AI駆動開発ガイド/preview
"""
from pathlib import Path
import json
import sys

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dads_theme import *  # noqa: F403

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "AI駆動開発ガイド"
OUT.mkdir(parents=True, exist_ok=True)
TITLE = "AI駆動開発ガイド"
DATE = "2026年9月8日"
W = 1152

SOURCES = {
    1: ("Anthropic｜Best practices for Claude Code",
        "https://code.claude.com/docs/en/best-practices",
        "タスクの具体化・文脈・実行による検証"),
    2: ("GitHub｜Application card: GitHub Copilot Agents",
        "https://docs.github.com/en/copilot/responsible-use/agents",
        "生成コードとAIレビューの限界・人の責任"),
    3: ("OWASP｜LLM01:2025 Prompt Injection",
        "https://genai.owasp.org/llmrisk/llm01-prompt-injection/",
        "外部コンテンツを介した指示の混入"),
    4: ("OWASP｜LLM06:2025 Excessive Agency",
        "https://genai.owasp.org/llmrisk/llm062025-excessive-agency/",
        "権限の最小化・重要操作の制御"),
    5: ("DORA｜Software delivery performance metrics",
        "https://dora.dev/guides/dora-metrics/",
        "変更の流れと不安定性を測る観点"),
}

prs = Presentation()
prs.slide_width, prs.slide_height = px(CANVAS_W), px(CANVAS_H)
prs.core_properties.title = TITLE
prs.core_properties.subject = "AIを使う開発チームのための実践ガイド"
prs.core_properties.author = ""
prs.core_properties.keywords = "AI駆動開発, 生成AI, ソフトウェア開発, DADS"
prs.core_properties.comments = "DADSの基本デザインをPowerPointへ写像。公式の監修資料ではありません。"
manifest = []
color_checks = set()


def t(s, x, y, w, h, text, style="Std-20N-150", color=INK,
      bg=WHITE, align="l", anchor="t"):
    """改行を個別段落にして、PPTXと近似プレビューを一致させる。"""
    assert contrast(color, bg) >= 4.5, (text, color, bg)
    color_checks.add((color, bg, "text"))
    return block(s, x, y, w, h,
                 [dict(text=line, style=style, color=color, align=align)
                  for line in text.split("\n")], anchor=anchor)


def panel(s, x, y, w, h, fill=WHITE, line=BORDER, radius=0):
    if line:
        assert contrast(line, fill) >= 3.0
        color_checks.add((line, fill, "boundary"))
    return rect(s, x, y, w, h, fill=fill, line=line, radius=radius)


def note(s, text):
    s.notes_slide.notes_text_frame.text = text


def new(title, section, lead="", refs=(), notes="", dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, CANVAS_W, CANVAS_H, fill=TERTIARY if dark else WHITE)
    i = len(prs.slides)
    manifest.append(dict(number=i, title=title, section=section, notes=notes, refs=list(refs)))
    note(s, f"{i:02d}｜{title}\n\n{notes}" +
         ("\n\n参考資料（2026年9月8日参照）\n" + "\n".join(
             f"[{r}] {SOURCES[r][0]}\n{SOURCES[r][1]}" for r in refs) if refs else ""))
    if not dark:
        rect(s, 64, 52, 6, 24, fill=PRIMARY)
        t(s, 80, 52, 980, 24, section, "Oln-14B-100", PRIMARY, anchor="m")
        t(s, 64, 84, W, 48, title, "Std-32B-150")
        rule(s, 64, 140, W, 2, BORDER)
        if lead:
            t(s, 64, 152, W, 32, lead, "Std-18N-160", INK_SUB)
        rule(s, 64, 660, W, 1, BORDER_SOFT)
        t(s, 64, 674, 440, 20, TITLE, "Dns-14N-130", INK_MUTE)
        if refs:
            t(s, 520, 674, 560, 20,
              "参考 " + " ".join(f"[{r}]" for r in refs) + " ｜出典は p.20",
              "Dns-14N-130", INK_MUTE, align="r")
        t(s, 1120, 674, 96, 20, f"{i:02d} / 20", "Dns-14B-130", INK_MUTE, align="r")
    return s


def band(s, text, y=576, h=64):
    panel(s, 64, y, W, h, BG_KEY, None)
    t(s, 88, y, W-48, h, text, "Std-20B-150", TERTIARY, BG_KEY, anchor="m")


def num(s, x, y, value, fill=PRIMARY):
    circle(s, x, y, 48, fill)
    t(s, x, y, 48, 48, value, "Std-20B-150", WHITE, fill, align="c", anchor="m")


def table(s, y, widths, labels, rows, row_h=72):
    panel(s, 64, y, W, 48, BG_KEY, None)
    x = 64
    for w, label in zip(widths, labels):
        t(s, x+16, y, w-32, 48, label, "Std-18B-160", PRIMARY, BG_KEY, anchor="m")
        x += w
    rule(s, 64, y+48, W, 2, PRIMARY)
    yy = y+50
    for row in rows:
        x = 64
        for j, (w, cell) in enumerate(zip(widths, row)):
            t(s, x+16, yy+4, w-32, row_h-8, cell,
              "Std-20B-150" if j == 0 else "Std-20N-150",
              INK if j == 0 else INK_SUB, anchor="m")
            x += w
        rule(s, 64, yy+row_h, W, 1, BORDER)
        yy += row_h+1
    return yy


# 01 ── 表紙
s = new(TITLE, "表紙", dark=True, notes=(
    "対象は、生成AIを開発業務に取り入れたいエンジニア、テックリード、開発マネージャー。"
    "30〜40分程度の勉強会、または作業時に参照するガイドを想定。製品の機能比較ではなく、"
    "小さな仕事を定義し、実行結果を確認する共通の進め方を扱う。"
    "サンプルの仕様、チェックリスト、4週間の導入案は本資料の提案例。"))
rect(s, 0, 0, 8, 720, fill=SECONDARY)
t(s, 80, 76, 720, 30, "AI-DRIVEN DEVELOPMENT / PRACTICAL GUIDE",
  "Oln-16B-100", WHITE, TERTIARY)
t(s, 80, 200, 760, 100, TITLE, "Dsp-64B-140", WHITE, TERTIARY)
rule(s, 80, 336, 120, 4, WHITE)
t(s, 80, 376, 720, 100, "小さく任せ、実行して確かめる。\n仕様からレビューまでの実践ガイド",
  "Std-28B-150", WHITE, TERTIARY)
for i, (title, sub) in enumerate([
    ("01  定義する", "目的・制約・受け入れ条件"),
    ("02  つくる", "調査・実装・テスト"),
    ("03  確かめる", "証拠・レビュー・改善"),
]):
    y = 184+i*128
    panel(s, 896, y, 320, 104, TERTIARY, WHITE)
    t(s, 920, y+16, 272, 36, title, "Std-24B-150", WHITE, TERTIARY)
    t(s, 920, y+58, 272, 30, sub, "Std-18N-160", WHITE, TERTIARY)
rule(s, 80, 620, 1136, 1, WHITE)
t(s, 80, 644, 850, 28, "開発チーム向け  /  入門からチーム運用まで", "Std-18N-160", WHITE, TERTIARY)
t(s, 952, 644, 264, 28, DATE, "Std-18N-160", WHITE, TERTIARY, align="r")

# 02 ── 読み方
s = new("このガイドで、最初の1件を完了できる", "GUIDE MAP",
        "全体像をつかみ、テンプレートを使い、確認項目に沿って実践する。",
        notes="最初からすべてを自動化しない。まずは自分で正誤を判断できる小さなタスクを選ぶ。"
        "p.08は依頼時、p.13はレビュー時、p.19は作業の前後に参照できる。")
for i, (head, body, pages) in enumerate([
    ("理解する", "人とAIの役割、任せる仕事、基本の流れ", "03—05"),
    ("依頼する", "実行環境、渡す情報、依頼文の型", "06—08"),
    ("つくって確かめる", "仕様例、実装、テスト、調査、レビュー", "09—14"),
    ("チームに定着させる", "運用ルール、効果測定、導入計画", "15—19"),
]):
    y = 204+i*100
    num(s, 64, y+18, f"{i+1:02}")
    t(s, 136, y+12, 280, 40, head, "Std-24B-150")
    t(s, 440, y+18, 624, 40, body, "Std-20N-150", INK_SUB)
    t(s, 1080, y+16, 136, 40, pages, "Std-24B-150", PRIMARY, align="r")
    rule(s, 136, y+86, 1080, 1, BORDER)

# 03 ── 定義と役割
s = new("人が目的を決め、AIと検証のループを回す", "01 / 理解する",
        "本資料では、AIを要件整理・設計・実装・検証に組み込む開発の進め方を「AI駆動開発」と呼ぶ。",
        refs=(2,), notes="AIが出力した説明やコードは、そのまま正しいという証拠にはならない。"
        "人が期待する挙動と制約を定め、AIは調査や作成を補助する。テストや実行結果を確認し、"
        "採用とリリースの責任はチームが持つ。役割は提案であり、案件のリスクに合わせて調整する。")
for i, (role, title, a, b, c) in enumerate([
    ("人", "目的と判断を持つ", "何を解決するか決める", "受け入れ条件を定める", "差分を理解し採用する"),
    ("AI", "調べて形にする", "既存コードを調査する", "設計案と変更をつくる", "テスト案・説明をつくる"),
    ("実行環境", "結果を返す", "テストとビルドを実行", "静的解析で問題を検出", "画面やログで挙動を確認"),
]):
    x = 64+i*392
    panel(s, x, 216, 368, 328)
    panel(s, x, 216, 368, 56, PRIMARY, None)
    t(s, x+24, 216, 320, 56, role, "Std-22B-150", WHITE, PRIMARY, anchor="m")
    t(s, x+24, 296, 320, 42, title, "Std-24B-150")
    for j, text in enumerate([a, b, c]):
        t(s, x+24, 360+j*48, 320, 40, text, "Std-20N-150", INK_SUB)
band(s, "完了の条件は、期待した挙動を確認でき、変更の責任者が説明できること。")

# 04 ── タスク選定
s = new("最初は、正誤を確かめやすい仕事を選ぶ", "01 / 理解する",
        "選定の3条件：範囲が小さい・期待結果が明確・元に戻せる。",
        notes="分類は本資料の導入提案。ドキュメントも正解の出典が必要で、無条件に低リスクではない。"
        "認証やDB移行もAIを使えるが、最初の試行対象にはしない。調査やテスト観点の整理に限定し、"
        "設計や実行計画を担当者が確認する。")
table(s, 208, [240, 468, 444], ["候補", "任せる作業の例", "完了を判断する証拠"], [
    ["単体テストの追加", "既存関数の境界値を追加", "期待値の根拠＋テスト結果"],
    ["小さな不具合修正", "再現条件が明確なバグを修正", "修正前に失敗、修正後に成功"],
    ["説明文の更新", "コードとREADMEの差を修正", "参照元と手順の実行結果"],
    ["認証・DB移行", "まずは影響調査と案の比較", "担当者の設計確認＋移行検証"],
], row_h=68)
band(s, "タスクが大きいときは「調査」「仕様の整理」「1つの変更」に分ける。")

# 05 ── 開発フロー
s = new("1つの変更を、5つの工程で進める", "01 / 理解する",
        "各工程で残す成果物を決めると、会話の途中からでも作業を再開できる。",
        refs=(1,2), notes="本資料の標準フロー案。仕様、計画、差分、実行記録、レビュー結果を残す。"
        "仕様が曖昧なら仕様へ、検証で問題が出たら計画や実装へ戻る。"
        "明確で軽微な修正では計画を一文で済ませるなど、手順の重さを調整する。")
steps = [("仕様", "何を満たすか", "受け入れ条件"), ("計画", "どこを変えるか", "対象と変更順"),
         ("実装", "小さく変更する", "コードの差分"), ("検証", "実際に動かす", "実行結果・画面"),
         ("レビュー", "採用を判断する", "確認済みのPR")]
for i, (head, desc, artifact) in enumerate(steps):
    x = 64+i*236
    panel(s, x, 232, 208, 256)
    panel(s, x, 232, 208, 52, PRIMARY if i != 4 else TERTIARY, None)
    t(s, x, 232, 208, 52, f"{i+1}  {head}", "Std-24B-150", WHITE,
      PRIMARY if i != 4 else TERTIARY, align="c", anchor="m")
    t(s, x+16, 312, 176, 64, desc, "Std-20B-150", align="c")
    t(s, x+16, 400, 176, 64, artifact, "Std-18N-160", INK_SUB, align="c")
    if i < 4:
        arrow_right(s, x+222, 352, 16, PRIMARY)
rule(s, 160, 528, 952, 2, PRIMARY)
rect(s, 1111, 488, 2, 42, fill=PRIMARY)
rect(s, 159, 508, 2, 22, fill=PRIMARY)
arrow = arrow_right(s, 160, 500, 16, PRIMARY)
arrow.rotation = 0
t(s, 304, 540, 680, 32, "未達・不明点があれば、必要な工程へ戻る", "Std-20B-150", PRIMARY, align="c")
t(s, 64, 608, W, 32, "PR（プルリクエスト）：変更内容を共有し、レビューして取り込むための単位。", "Std-18N-160", INK_SUB)

# 06 ── 実行環境
s = new("使い方に合わせて、実行できる範囲を決める", "02 / 依頼する",
        "組織で利用できるツールを選び、編集・コマンド実行・外部接続の範囲を確認する。",
        refs=(2,4), notes="分類は概念上のもの。実際の製品は複数の使い方を併せ持ち、機能や契約で権限が異なる。"
        "最初に入力データの扱い、認証、利用可能なモデル、アクセス範囲、利用料とログを確認する。"
        "環境構築と基準となるテストを済ませておけば、AIの変更による不具合を切り分けやすい。")
for i, (head, use, check) in enumerate([
    ("補完", "書いている箇所の\n候補を得る", "候補の採用前に\n挙動と依存関係を読む"),
    ("対話", "仕様・設計の相談や\nコードの説明を得る", "回答の参照元と\n前提条件を確認する"),
    ("エージェント", "複数ファイルの変更や\nテスト実行を任せる", "編集範囲・実行権限・\n外部接続を制御する"),
]):
    x = 64+i*392
    panel(s, x, 208, 368, 316)
    t(s, x+24, 232, 320, 48, head, "Std-28B-150", PRIMARY)
    t(s, x+24, 300, 320, 68, use, "Std-22B-150")
    rule(s, x+24, 392, 320, 1, BORDER)
    t(s, x+24, 412, 320, 88, check, "Std-20N-150", INK_SUB)
band(s, "開始前に、作業ブランチ・テスト環境・既存テストの結果をそろえる。", y=556, h=80)

# 07 ── 文脈
s = new("必要な情報を、参照先と優先順位で渡す", "02 / 依頼する",
        "判断に使う仕様、対象ファイル、実行方法を短くそろえる。",
        refs=(1,), notes="コンテキストとはAIが判断に使える情報。仕様、関連コード、コマンド、直近のエラーなどが含まれる。"
        "すべてのファイルや会話を詰め込むと必要情報が埋もれるため、関連範囲と参照先を示す。"
        "継続作業では、決定事項、変更済みファイル、検証結果、未解決点を短い引き継ぎメモに残す。")
for i, (head, body) in enumerate([
    ("タスク仕様", "目的・入出力・受け入れ条件・対象外"),
    ("リポジトリ", "関連ファイル・既存の実装例・設計判断"),
    ("実行環境", "セットアップ・テストコマンド・再現ログ"),
]):
    y = 208+i*120
    num(s, 64, y+24, f"{i+1:02}")
    t(s, 136, y+12, 640, 40, head, "Std-24B-150")
    t(s, 136, y+60, 640, 40, body, "Std-20N-150", INK_SUB)
    rule(s, 136, y+108, 616, 1, BORDER)
panel(s, 808, 208, 408, 348, BG_KEY, None)
t(s, 840, 236, 344, 40, "参照のルール", "Std-24B-150", TERTIARY, BG_KEY)
t(s, 840, 300, 344, 216,
  "最新の仕様を明示する\n該当ファイルを絞る\nバージョンを伝える\n不明点は推測と区別する\n決定事項を更新する",
  "Std-22B-150", TERTIARY, BG_KEY)
band(s, "長い作業では「決定事項・変更・検証済み・未解決」の4点を引き継ぐ。", y=584, h=56)

# 08 ── 依頼文
s = new("依頼文は、目的・範囲・完了条件まで書く", "02 / 依頼する",
        "コピーして使うテンプレート例：次ページの normalize_title を実装する。",
        refs=(1,), notes="以下はそのまま依頼文として使える例。\n"
        "目的：申請メモのタイトルを保存前に整える。\n"
        "対象：src/title.py と tests/test_title.py。\n"
        "仕様：p.09の受け入れ条件に従う。\n"
        "参照：既存の入力チェックとテストの書き方。\n"
        "制約：公開APIの変更と依存追加は行わない。\n"
        "進め方：関連箇所を確認し、短い計画を示して実装する。仕様上の不明点は実装前に確認する。\n"
        "完了条件：対象テストと既存テストを実行。差分の要点、コマンドと結果、未検証事項を報告する。\n"
        "ファイルパスやコマンドはサンプルであり、実際のリポジトリに合わせて変更する。")
panel(s, 64, 204, 792, 436, GRAY[50], None)
fields = [
    ("目的", "申請メモのタイトルを、保存前に整える。"),
    ("対象・仕様", "src/title.py とテスト。条件は p.09 に従う。"),
    ("参照・制約", "既存の実装に合わせる。公開APIと依存は維持。"),
    ("進め方", "関連箇所を確認し、短い計画を示して実装。"),
    ("完了条件", "対象テストと既存テストを実行して確認。"),
    ("報告形式", "差分の要点／実行コマンドと結果／未検証事項。"),
]
for i, (key, value) in enumerate(fields):
    y = 224+i*66
    t(s, 88, y, 160, 30, key, "Std-18B-160", PRIMARY, GRAY[50])
    t(s, 264, y, 568, 54, value, "Std-20N-150", INK, GRAY[50])
panel(s, 888, 204, 328, 436, BG_KEY, None)
t(s, 912, 228, 280, 84, "「できた」を\n判定できる依頼に", "Std-24B-150", TERTIARY, BG_KEY)
for i, value in enumerate(["何が変わるか", "どこまで変えるか", "何を実行するか", "何を提出するか"]):
    t(s, 912, 336+i*64, 280, 44, value, "Std-22B-150", TERTIARY, BG_KEY)

# 09 ── 仕様例
s = new("受け入れ条件は、入力と期待結果で決める", "03 / つくって確かめる",
        "題材：normalize_title(value) — 申請メモのタイトルを保存前に整える関数。",
        notes="本資料のオリジナル仕様例。文字数はPythonのlenで数えるUnicodeコードポイント数とする。"
        "画面上の見た目の文字数とは異なる場合がある。前後の空白はstr.strip()で除去し、内部の空白は維持する。"
        "処理順は型確認、前後の空白除去、空文字確認、長さ確認、戻り値の順。DB保存や画面表示、"
        "Unicode正規化はこのタスクの範囲に含めない。仕様判断は実装に先立って確定させる。")
table(s, 204, [268, 408, 476], ["観点", "受け入れ条件", "入力 → 期待結果"], [
    ["前後の空白", "前後を除去し、内部は保つ", '「  申請 メモ  」→「申請 メモ」'],
    ["必須入力", "空文字・空白のみはエラー", '「   」→ ValueError'],
    ["最大長", "除去後の長さは1〜50", "50文字 → 成功／51文字 → エラー"],
    ["型", "文字列以外はエラー", "123、None → TypeError"],
], row_h=72)
band(s, "数え方も仕様にする：ここでは Python の len() を使う。DB保存・画面表示は対象外。", y=576, h=64)

# 10 ── 実装
s = new("変更は、レビューできる小ささに保つ", "03 / つくって確かめる",
        "受け入れ条件に必要な差分をつくり、途中で変更範囲を確かめる。",
        notes="サンプルはPythonを想定した実装の順序。型ヒントだけで実行時の型チェックは行われないため、"
        "文字列以外を拒否する条件を明示する。タイトルの正規化と保存処理を混ぜない。"
        "既存コードを読むことで、例外の種類や配置に関するプロジェクトの方針と整合させる。"
        "変更後はdiffを確認し、不要な整形や依存追加、無関係な修正が入っていないかを見る。")
for i, (title, desc) in enumerate([
    ("既存の実装を読む", "同じ責務の処理・例外・テストを確認"),
    ("変更範囲を示す", "関数とテストの対象ファイルを列挙"),
    ("1つの仕様を実装する", "型 → 空白除去 → 必須 → 最大長の順"),
    ("差分と実行結果を確認", "無関係な修正がないかも読む"),
]):
    y = 208+i*100
    num(s, 64, y+8, f"{i+1:02}")
    t(s, 136, y, 592, 40, title, "Std-24B-150")
    t(s, 136, y+44, 616, 36, desc, "Std-20N-150", INK_SUB)
panel(s, 816, 208, 400, 404, TERTIARY, None)
t(s, 844, 240, 344, 48, "変更の単位", "Std-24B-150", WHITE, TERTIARY)
t(s, 844, 324, 344, 140, "1つの目的\n1つの確認可能な結果\n1つの説明できる差分", "Std-24B-150", WHITE, TERTIARY)
t(s, 844, 520, 344, 64, "複雑なら、先に分割する。", "Std-20N-150", WHITE, TERTIARY)

# 11 ── テスト
s = new("テストの期待値は、受け入れ条件から決める", "03 / つくって確かめる",
        "AIが書いた実装とテストが、同じ誤解を共有する可能性を考える。",
        refs=(2,), notes="テスト数を増やすだけでなく、仕様に根拠のある期待値を確認する。"
        "境界値の前後、空文字、空白、型違いを用意する。不具合修正では修正前のコードで再現テストが失敗し、"
        "修正後に成功することを確認する。新規関数ではまずテストを実行し、未実装状態で成功していないことを確認する。"
        "実行コマンドと結果を添え、実行できないテストは未検証として記録する。")
table(s, 204, [228, 504, 420], ["種類", "テスト入力の例", "期待する結果"], [
    ["正常系", '「  申請 メモ  」', '「申請 メモ」'],
    ["境界値", '「あ」× 1、50、51', "1・50は成功、51は ValueError"],
    ["異常系", '空文字、半角空白のみ', "ValueError"],
    ["型違い", "123、None", "TypeError"],
], row_h=66)
panel(s, 64, 552, W, 88, BG_KEY, None)
t(s, 88, 564, 1104, 32, "提出する証拠：対象テスト＋既存テストのコマンド、結果、未実行の範囲。",
  "Std-20B-150", TERTIARY, BG_KEY)
t(s, 88, 602, 1104, 28, "修正タスクでは、再現テストが「修正前に失敗 → 修正後に成功」することも確認する。",
  "Std-18N-160", TERTIARY, BG_KEY)

# 12 ── デバッグ
s = new("行き詰まったら、再現条件に戻る", "03 / つくって確かめる",
        "修正を重ねる前に、観測した事実と原因の仮説を分ける。",
        notes="たとえば『50文字のはずなのにエラーになる』場合、計数方法、空白除去のタイミング、"
        "画面とサーバーの差などを仮説として挙げる。ログは必要範囲に限定し、機密情報は除く。"
        "AIには各仮説の確認方法を求め、1つずつ検証する。根拠のない試行が続く場合は最後に確認できた状態へ戻り、"
        "再現手順、関連ファイル、試したことを整理する。")
panel(s, 64, 208, 552, 356, GRAY[50], None)
t(s, 88, 232, 504, 40, "AIに渡す調査メモ", "Std-24B-150", INK, GRAY[50])
for i, text in enumerate([
    "期待：50文字のタイトルを受け付ける",
    "実際：入力によって ValueError になる",
    "条件：入力値・操作手順・バージョン",
    "証拠：エラー全文・関連ログ・差分",
    "依頼：仮説と確認方法を提示する",
]):
    t(s, 88, 296+i*48, 504, 40, text, "Std-20N-150", INK_SUB, GRAY[50])
for i, (title, desc) in enumerate([
    ("再現する", "最小の入力と手順に絞る"),
    ("仮説を確かめる", "1回に1つの原因を検証する"),
    ("回帰を防ぐ", "再現テストを残して修正する"),
]):
    y = 216+i*116
    num(s, 664, y, f"{i+1:02}")
    t(s, 736, y, 480, 42, title, "Std-24B-150")
    t(s, 736, y+48, 480, 42, desc, "Std-20N-150", INK_SUB)
band(s, "説明が変わるだけで証拠が増えないときは、差分を広げず状況を整理する。", y=584, h=56)

# 13 ── レビュー
s = new("レビューでは、差分と証拠を一緒に読む", "03 / つくって確かめる",
        "AIレビューの指摘も検証し、採用する変更は担当者が説明できる状態にする。",
        refs=(2,), notes="AIレビューには見落としや誤検出がある。指摘箇所と再現可能性を確認してから変更する。"
        "PRには目的、主要な差分、検証、未検証事項、影響と戻し方を記載する。"
        "CI（継続的インテグレーション）は変更時にテストや解析を自動実行する仕組み。"
        "画面操作や業務上の妥当性など、自動テストで十分確認できない項目は担当者が確認する。")
panel(s, 64, 208, 552, 432)
t(s, 88, 232, 504, 42, "レビューの5観点", "Std-24B-150", PRIMARY)
for i, (head, body) in enumerate([
    ("仕様", "受け入れ条件と対象外を守る"),
    ("構造", "既存設計に沿い、差分が適切"),
    ("安全", "権限・入力・機密情報を確認"),
    ("証拠", "テストと解析の実行結果がある"),
    ("運用", "ログ・影響・戻し方を説明できる"),
]):
    y = 296+i*60
    t(s, 88, y, 88, 40, head, "Std-20B-150")
    t(s, 192, y, 400, 52, body, "Std-20N-150", INK_SUB)
panel(s, 648, 208, 568, 432, BG_KEY, None)
t(s, 676, 232, 512, 42, "PR説明文の型", "Std-24B-150", TERTIARY, BG_KEY)
for i, (head, desc) in enumerate([
    ("目的・変更", "何が、どのように変わるか"),
    ("検証結果", "実行コマンドと結果、必要なら画面"),
    ("未検証・影響", "確認できていない範囲と利用者への影響"),
    ("戻し方", "切り戻し手順と注意点"),
]):
    y = 296+i*80
    t(s, 676, y, 512, 30, head, "Std-18B-160", TERTIARY, BG_KEY)
    t(s, 676, y+32, 512, 40, desc, "Std-20N-150", INK, BG_KEY)

# 14 ── 安全な利用
s = new("データと実行権限を、必要な範囲に絞る", "03 / つくって確かめる",
        "ツールの管理設定と実行環境で制御し、重要な操作には確認を入れる。",
        refs=(3,4), notes="OWASPは外部コンテンツ経由のプロンプトインジェクションや過剰な権限によるリスクを整理している。"
        "コードコメントやWeb文書に混入した指示を、ユーザーや組織の指示として扱わない。"
        "プロンプトへの注意書きだけでは防御を保証できないため、アクセス権、サンドボックス、接続先制限、ログを組み合わせる。"
        "機密情報や個人情報の入力可否、保存・学習利用の設定は組織の契約とルールに従って確認する。"
        "本番への変更、データ削除、外部送信は、具体的な対象と影響を示して権限を持つ担当者が判断する。")
for i, (head, a, b, c) in enumerate([
    ("入力データ", "承認済みの環境を使う", "秘密鍵や実データを除く", "保持・学習利用の設定を確認"),
    ("実行権限", "作業範囲と接続先を絞る", "本番の認証情報を分離する", "操作と実行結果を記録する"),
    ("外部コンテンツ", "文書やコメントを資料扱い", "混入した命令を採用しない", "要求外の外部送信を止める"),
]):
    x = 64+i*392
    panel(s, x, 208, 368, 328)
    t(s, x+24, 236, 320, 48, head, "Std-28B-150", PRIMARY)
    for j, text in enumerate([a, b, c]):
        t(s, x+24, 316+j*66, 320, 64, text, "Std-20N-150", INK_SUB)
band(s, "本番変更・データ削除・外部送信は、対象と影響を示して担当者が判断する。", y=568, h=72)

# 15 ── チーム標準
s = new("チームで共有するルールは、短く保つ", "04 / チームに定着させる",
        "人が読む開発ガイドと、AIが参照する指示を、同じ運用ルールにそろえる。",
        refs=(1,), notes="本資料のチーム標準案。共通ルールには利用可能なデータ、実行範囲、開発と検証のコマンド、"
        "レビューの基準を記載する。AI用指示ファイルの名前や読込仕様は使用するツールの公式文書で確認する。"
        "プロンプト上の禁止事項だけでは権限制御にならないため、設定やCIでも強制する。"
        "環境や規約を変更した担当者がルールを更新し、導入初期は週1回の振り返りで失敗例を反映する。")
table(s, 204, [256, 608, 288], ["決めること", "最低限そろえる内容", "主な管理者の例"], [
    ["利用範囲", "ツール・データ・外部接続・重要操作", "管理者／セキュリティ"],
    ["開発の約束", "ディレクトリ・設計方針・依存の扱い", "テックリード"],
    ["検証の手順", "セットアップ・テスト・解析コマンド", "開発チーム"],
    ["完了の基準", "必須チェック・レビュー担当・証拠", "開発責任者"],
], row_h=70)
band(s, "ルールの管理者と更新タイミングを決め、コードや環境の変更に追随させる。", y=576, h=64)

# 16 ── アンチパターン
s = new("よくある失敗は、作業の進め方から直す", "04 / チームに定着させる",
        "出力だけを直し続けず、仕様・情報・検証・分割のどこに原因があるかを見る。",
        notes="本資料で整理した典型例。仕様の曖昧さに対して出力の再生成だけを繰り返すと、"
        "毎回違う仮定が入る。CIを通すためだけにテストの期待値を書き換えると、本来の不具合が隠れる。"
        "AIの採用量を増やすこと自体を目的にせず、レビュー可能性と利用者の期待する挙動を優先する。")
table(s, 204, [360, 336, 456], ["起きがちなこと", "考えられる原因", "次に変える行動"], [
    ["生成のたびに方針が変わる", "仕様・判断が未確定", "決定事項と対象外を明記"],
    ["大量の差分を読み切れない", "タスクの単位が大きい", "1つの目的ごとに分ける"],
    ["成功の説明だけで完了する", "実行可能な確認がない", "実行コマンドと結果を求める"],
    ["テストを緩めて成功させる", "期待値の根拠が曖昧", "元の仕様と再現条件に戻る"],
    ["同じ説明を繰り返す", "参照情報が分散・古い", "短い共通ルールを更新する"],
], row_h=70)

# 17 ── 効果測定
s = new("効果は、速さ・品質・負担を一緒に測る", "04 / チームに定着させる",
        "導入前と同種の仕事を比べ、レビューと手戻りを含めた全体の変化を見る。",
        refs=(5,), notes="DORAはソフトウェア変更の流れと不安定性を測る指標を提供する。本スライドは、"
        "その視点を参考にした小規模導入向けの測定案であり、DORAの指標セットをそのまま列挙したものではない。"
        "変更リードタイムはコミットから本番反映までとし、別に着手からレビュー完了までを測るなら名称と起終点を分ける。"
        "同じチーム、同程度の難易度、同じ集計期間で比較する。小標本や自己申告だけで因果や倍率を断定しない。"
        "AI利用料に加えてレビュー、修正、環境整備の工数も記録する。個人のランキングには使わない。")
for i, (head, metric, body) in enumerate([
    ("速さ", "変更リードタイム", "コミットから本番反映まで\nレビュー待ち時間も記録"),
    ("品質", "不具合・手戻り", "本番障害や差し戻しの件数\n原因と修正工数を記録"),
    ("負担", "確認時間・総コスト", "レビューと修正の時間\nAI利用料と運用工数を記録"),
]):
    x = 64+i*392
    panel(s, x, 216, 368, 304)
    t(s, x+24, 240, 320, 48, head, "Std-28B-150", PRIMARY)
    t(s, x+24, 320, 320, 40, metric, "Std-24B-150")
    t(s, x+24, 392, 320, 92, body, "Std-20N-150", INK_SUB)
panel(s, 64, 552, W, 88, BG_KEY, None)
t(s, 88, 564, 1104, 32, "比較の条件：同じチーム・同種のタスク・同じ集計期間。", "Std-20B-150", TERTIARY, BG_KEY)
t(s, 88, 604, 1104, 28, "生成行数や採用率だけで成功とせず、品質の悪化とレビュー負荷も確認する。", "Std-18N-160", TERTIARY, BG_KEY)

# 18 ── 導入計画
s = new("4週間で、小さく試して継続を判断する", "04 / チームに定着させる",
        "導入計画の例：1チーム・1種類のタスクから始める。期間と規模は調整する。",
        notes="本資料の導入提案であり、4週間で効果を保証するものではない。"
        "1週目に過去の同種タスクから比較用データを取り、タスク、担当者、利用範囲、受け入れ条件を決める。"
        "2週目は少数のタスクで実践し、3週目は成功例と失敗例を共通ルールに反映する。"
        "4週目に速さ、品質、負担を比較する。件数が足りない場合は結論を急がず試行期間を延ばす。"
        "改善がなくても、対象を変える、検証を整える、利用を見直すという判断につなげる。")
weeks = [
    ("WEEK 1", "準備", "対象と担当を決める\n環境・基準値を用意", "成果物：試行計画"),
    ("WEEK 2", "実践", "小さなタスクで試す\n時間と失敗を記録する", "成果物：実践ログ"),
    ("WEEK 3", "改善", "依頼と検証を見直す\nよい例を共有する", "成果物：共通の依頼文"),
    ("WEEK 4", "判断", "同種タスクと比べる\n継続・拡大・見直し", "成果物：判断と次の対象"),
]
for i, (week, head, body, deliverable) in enumerate(weeks):
    x = 64+i*296
    panel(s, x, 216, 264, 332)
    panel(s, x, 216, 264, 48, PRIMARY, None)
    t(s, x+24, 216, 216, 48, week, "Oln-16B-100", WHITE, PRIMARY, anchor="m")
    t(s, x+24, 284, 216, 48, head, "Std-28B-150")
    t(s, x+24, 352, 216, 112, body, "Std-20N-150", INK_SUB)
    t(s, x+24, 480, 216, 52, deliverable, "Std-18B-160", PRIMARY)
    if i < 3:
        arrow_right(s, x+280, 388, 16, PRIMARY)
band(s, "拡大は、品質・負担を許容範囲に保ち、チームで繰り返し実践できてから。", y=580, h=60)

# 19 ── チェックリスト
s = new("明日の1件は、このチェックから始める", "04 / チームに定着させる",
        "開始時と完了時に確認する、持ち帰り用チェックリスト。",
        notes="最初の対象として、再現できる小さな不具合や既存関数のテスト追加を選ぶ。"
        "開始前に目的、範囲、受け入れ条件、参照情報、環境を確認し、完了時には差分、実行結果、"
        "未検証事項、担当者の確認をそろえる。チェックが埋まらない項目を、次の改善対象にする。"
        "すべてを自動化する必要はなく、自分たちで扱える範囲から学習を積み重ねる。")
for x, heading, items in [
    (64, "開始前", ["目的と対象外を説明できる", "受け入れ条件が具体的", "関連ファイルと実行方法がある", "データと実行権限を確認した"]),
    (656, "完了前", ["差分を読み、意図を説明できる", "テスト・解析の実行結果がある", "未検証の範囲を明記した", "レビュー担当者が確認した"]),
]:
    panel(s, x, 208, 560, 328)
    t(s, x+24, 232, 512, 48, heading, "Std-28B-150", PRIMARY)
    for i, value in enumerate(items):
        y = 304+i*52
        panel(s, x+28, y+7, 20, 20, WHITE, PRIMARY)
        t(s, x+64, y, 472, 40, value, "Std-22B-150")
panel(s, 64, 568, W, 72, TERTIARY, None)
t(s, 88, 568, 1104, 72, "小さく任せる。結果を確かめる。学びをチームに残す。",
  "Std-28B-150", WHITE, TERTIARY, align="c", anchor="m")

# 20 ── 出典
s = new("参考資料・出典", "REFERENCES",
        "2026年9月8日参照。一般化した実践手順と、独自の仕様例・導入案を含みます。",
        notes="各URLはスライド上でクリック可能。各ページの参考番号と対応する。"
        "製品の操作方法や契約条件は随時変わるため、導入時に組織の設定と公式文書を確認する。"
        "デザインはDADSの基本トークンをPowerPointに写像したもので、公式の作成・監修資料ではない。"
        "PPTXは編集可能なテキストと図形で構成する。フォントはNoto Sans JPを指定。")
for i, (key, (name, url, purpose)) in enumerate(SOURCES.items()):
    y = 200+i*66
    t(s, 64, y, 64, 30, f"[{key}]", "Std-18B-160", PRIMARY)
    t(s, 128, y, 744, 30, name, "Std-18B-160")
    tb = t(s, 128, y+30, 760, 28, url, "Dns-16N-130", PRIMARY)
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    tb.text_frame.paragraphs[0].runs[0].font.underline = True
    t(s, 912, y+4, 304, 56, purpose, "Std-16N-170", INK_SUB)
rule(s, 64, 544, W, 1, BORDER)
t(s, 64, 562, W, 50,
  "出典：デジタル庁デザインシステムウェブサイト https://design.digital.go.jp/dads/ のコンテンツを加工して作成",
  "Std-16N-170", INK_SUB)
t(s, 64, 610, W, 28, "本資料はデジタル庁が作成・監修したものではありません。",
  "Std-16N-170", INK_SUB)


assert len(prs.slides) == 20
for i, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (i, shape.name)
        assert shape.left+shape.width <= prs.slide_width+px(0.5), (i, shape.name)
        assert shape.top+shape.height <= prs.slide_height+px(0.5), (i, shape.name)
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert run.font.size.pt >= 10.5, (i, run.text)
    assert slide.has_notes_slide

path = OUT / "AI駆動開発ガイド_DADS.pptx"
prs.save(path)

# 内容と出典をテキストでも参照できるよう保存する。
md = [f"# {TITLE}", "", f"作成日：{DATE} / 全20枚 / 16:9", "",
      "対象：開発者・テックリード・開発マネージャー。30〜40分程度の勉強会を想定。", ""]
for entry, slide in zip(manifest, prs.slides):
    md.extend([f"## {entry['number']:02d} {entry['title']}", ""])
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            if shape.top < px(660):
                md.extend([shape.text, ""])
    md.extend(["### 発表者ノート", "", entry["notes"], ""])
    for ref in entry["refs"]:
        md.append(f"参考 [{ref}] [{SOURCES[ref][0]}]({SOURCES[ref][1]})")
    md.append("")
(OUT / "AI駆動開発ガイド_原稿.md").write_text("\n".join(md), encoding="utf-8")
(OUT / "ai_guide_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
report = {
    "slides": len(prs.slides), "canvas_css_px": [1280, 720],
    "font": FONT, "minimum_font_css_px": 14,
    "shape_bounds": "pass", "speaker_notes": "20 / 20",
    "contrast_checks": [dict(fg=fg, bg=bg, role=role, ratio=round(contrast(fg, bg), 2))
                        for fg, bg, role in sorted(color_checks)],
    "render_check": "Run scripts/render_check.py separately.",
    "limitation": "Preview is an approximation; actual PowerPoint rendering is not verified.",
}
(OUT / "ai_guide_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Saved {path}\nSlides: {len(prs.slides)}; bounds/fonts/contrast: passed")
