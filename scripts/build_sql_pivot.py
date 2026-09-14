# -*- coding: utf-8 -*-
"""SQLで縦持ちを横持ちに変換してクロス集計する手順（アニメーション付き）を生成する。

実行: .venv/bin/python scripts/build_sql_pivot.py
検証: .venv/bin/python scripts/render_check.py out/SQLクロス集計/SQLクロス集計_DADS.pptx out/SQLクロス集計/preview
      .venv/bin/python scripts/dads_anim.py out/SQLクロス集計/SQLクロス集計_DADS.pptx

アニメーションはDADS（WCAG 2.2）の方針に合わせてフェードのみ・クリック進行とし、
最終状態（全要素表示）だけで内容が伝わる構成にする。プレビューPNGは最終状態を描画する。
"""
from pathlib import Path
import itertools
import json
import sys

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dads_theme import *  # noqa: F403
from dads_anim import grab, Timeline, code_line, code_block, CODE_BG

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "SQLクロス集計"
OUT.mkdir(parents=True, exist_ok=True)
PPTX = OUT / "SQLクロス集計_DADS.pptx"
TITLE = "SQLで縦持ちを横持ちへ ─ クロス集計の手順"
TOTAL = 13
W = CANVAS_W - MARGIN_X * 2

prs = Presentation()
prs.slide_width, prs.slide_height = px(CANVAS_W), px(CANVAS_H)
prs.core_properties.title = TITLE
prs.core_properties.subject = "CASE式とGROUP BYによるクロス集計（アニメーション付き）"
prs.core_properties.author = ""
prs.core_properties.keywords = "SQL, クロス集計, ピボット, 縦持ち, 横持ち, DADS"
prs.core_properties.comments = "DADSの基本デザインをPowerPointへ写像。公式の監修資料ではありません。"
manifest = []


# ---------------------------------------------------------------- 描画ヘルパ（dads_theme の薄いラッパ）
def t(s, x, y, w, h, text, style="Std-17N-170", color=INK, bg=WHITE,
      align="l", anchor="t", bullet=None):
    """改行を個別段落にする。背景色とのコントラストを検算する。"""
    assert contrast(color, bg) >= 4.5, (text, color, bg)
    return block(s, x, y, w, h,
                 [dict(text=ln, style=style, color=color, align=align, bullet=bullet)
                  for ln in text.split("\n")], anchor=anchor)


def panel(s, x, y, w, h, fill=WHITE, line=BORDER, radius=0, line_px=1, on=WHITE):
    if line:
        assert contrast(line, fill or on) >= 3.0, (line, fill or on)
    return rect(s, x, y, w, h, fill=fill, line=line, radius=radius, line_px=line_px)


def new(title, section, lead="", dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, CANVAS_W, CANVAS_H, fill=TERTIARY if dark else WHITE)
    i = len(prs.slides)
    manifest.append(dict(number=i, title=title, section=section, clicks=[]))
    if not dark:
        rect(s, MARGIN_X, 52, 6, 26, fill=PRIMARY)
        t(s, MARGIN_X + 16, 52, 900, 26, section, "Oln-14B-100", PRIMARY, anchor="m")
        t(s, MARGIN_X, 84, W, 48, title, "Std-32B-150")
        rule(s, MARGIN_X, 140, W, 2, BORDER)
        if lead:
            t(s, MARGIN_X, 152, W, 30, lead, "Std-17N-170", INK_SUB)
        rule(s, MARGIN_X, 660, W, 1, BORDER_SOFT)
        t(s, MARGIN_X, 672, 700, 20, TITLE, "Dns-14N-130", INK_MUTE)
        t(s, CANVAS_W - MARGIN_X - 200, 672, 200, 20, f"{i:02d} / {TOTAL}",
          "Dns-14B-130", INK_MUTE, align="r")
    return s


def band(s, text, y=568, h=64):
    panel(s, MARGIN_X, y, W, h, BG_KEY, None)
    t(s, MARGIN_X + 32, y, W - 64, h, text, "Std-20B-150", TERTIARY, BG_KEY, anchor="m")


def num(s, x, y, value, d=48, fill=PRIMARY):
    circle(s, x, y, d, fill)
    t(s, x, y, d, d, value, "Std-20B-150", WHITE, fill, align="c", anchor="m")


# ---------------------------------------------------------------- 表（横罫線ベース・セル単位で図形を返す）
def grid(s, x, y, widths, labels, rows, row_h=40, head_h=40, num_from=1):
    tw = sum(widths)
    xs = list(itertools.accumulate([x] + widths[:-1]))
    panel(s, x, y, tw, head_h, BG_KEY, None)
    for i, label in enumerate(labels):
        t(s, xs[i] + 12, y, widths[i] - 24, head_h, label, "Dns-14B-130", PRIMARY, BG_KEY,
          align="r" if i >= num_from else "l", anchor="m")
    rule(s, x, y + head_h, tw, 2, PRIMARY)
    cy = y + head_h + 2
    cells = []
    for r in rows:
        line = []
        for i, cell in enumerate(r):
            with grab(s) as g:
                if cell != "":
                    null = cell == "NULL"
                    t(s, xs[i] + 12, cy, widths[i] - 24, row_h, cell,
                      "Std-16B-170" if i == 0 else "Std-16N-170",
                      INK_MUTE if null else (INK if i == 0 else INK_SUB),
                      align="r" if i >= num_from else "l", anchor="m")
            line.append(g.shapes)
        rule(s, x, cy + row_h, tw, 1, BORDER_SOFT)
        cells.append(line)
        cy += row_h + 1
    return dict(cells=cells, xs=xs, widths=widths, top=y + head_h + 2, row_h=row_h,
                bottom=cy, w=tw)


def col(cells, c):
    return [shp for row in cells for shp in row[c]]


def row(cells, r):
    return [shp for cell in cells[r] for shp in cell]


def finish(tl):
    """アニメーションを書き込み、クリック手順を manifest に残す。"""
    manifest[-1]["clicks"] = tl.labels
    tl.apply()


# ---------------------------------------------------------------- サンプルデータ
LONG = [["東京", "4月", "120"], ["東京", "4月", "30"], ["東京", "5月", "140"],
        ["東京", "6月", "160"], ["大阪", "4月", "90"], ["大阪", "5月", "110"]]
MONTHS = ["4月", "5月", "6月"]


def spread(r):
    """CASE式の結果行（該当しない月は NULL）。"""
    return [r[0]] + [r[2] if r[1] == m else "NULL" for m in MONTHS]


WIDE = [["東京", "150", "140", "160"], ["大阪", "90", "110", "NULL"]]


# ================================================================ 01 表紙
s = new(TITLE, "表紙", dark=True)
rect(s, 0, 0, 8, CANVAS_H, fill=SECONDARY)
t(s, MARGIN_X + 24, 196, 900, 32, "SQL CROSS TABULATION", "Oln-16B-100", BLUE[300], TERTIARY)
t(s, MARGIN_X + 24, 240, 1080, 90, "縦持ちを横持ちへ", "Dsp-64B-140", WHITE, TERTIARY)
rect(s, MARGIN_X + 24, 360, 120, 4, fill=SECONDARY)
t(s, MARGIN_X + 24, 392, 1000, 80,
  "SQLでクロス集計する4つの手順\nCASE式・集計関数・GROUP BY の組み合わせ", "Std-24N-150", BLUE[100], TERTIARY)
rule(s, MARGIN_X + 24, 600, W - 48, 1, BLUE[1100])
t(s, MARGIN_X + 24, 620, 700, 24, "データベース勉強会資料｜クリックで順に表示します", "Dns-16N-130",
  BLUE[200], TERTIARY)
finish(Timeline(s))

# ================================================================ 02 ゴール
s = new("この資料のゴール", "はじめに", "縦に積まれたデータを、表として読み比べられる形にする。")
tl = Timeline(s)
goals = [("形の違いを説明できる",
          "縦持ちと横持ちの違いと、それぞれが向いている場面を理解する。", "p.3"),
         ("SQLで横持ちに変換できる",
          "CASE式・集計関数・GROUP BY を組み合わせて、どのDBでも動くクロス集計を書く。", "p.4〜9"),
         ("落とし穴を避けられる",
          "NULLと0の違い、COUNTの書き方、DBごとの専用構文（PIVOT・FILTER）を押さえる。", "p.10〜11")]
cw = (W - 32) / 3
for i, (head, body, ref) in enumerate(goals):
    x = MARGIN_X + i * (cw + 16)
    with grab(s) as g:
        panel(s, x, 208, cw, 312, WHITE, BORDER, RADIUS["m"])
        num(s, x + 24, 232, str(i + 1))
        t(s, x + 24, 300, cw - 48, 40, head, "Std-24B-150")
        t(s, x + 24, 352, cw - 48, 120, body, "Std-16N-170", INK_SUB)
        t(s, x + 24, 476, cw - 48, 24, "対応：" + ref, "Dns-14B-130", PRIMARY)
    tl.click(f"ゴール{i + 1}：{head}", g.shapes)
with grab(s) as g:
    band(s, "前提：SELECT・WHERE・GROUP BY の基本を知っていること", y=552, h=72)
tl.click("前提の表示", g.shapes)
finish(tl)

# ================================================================ 03 縦持ちと横持ち
s = new("縦持ちと横持ち", "01 ── 基本", "同じデータでも、行に積むか列に並べるかで形が変わる。")
tl = Timeline(s)
badge(s, MARGIN_X, 200, "縦持ち（ロング形式）", fill=BG_KEY, fg=PRIMARY, w=184)
grid(s, MARGIN_X, 240, [160, 160, 200], ["store（店舗）", "month（月）", "amount（売上）"], LONG,
     num_from=2)
t(s, MARGIN_X, 552, 520, 72,
  "1件の売上＝1行。月が増えても列は増えない\nDBへの保存・追記・絞り込みに向く", bullet="・")
with grab(s) as g:
    arrow_right(s, 640, 320, 28, PRIMARY)
    badge(s, 696, 200, "横持ち（ワイド形式）", fill=PRIMARY, fg=WHITE, w=184)
    grid(s, 696, 240, [140, 120, 120, 140], ["store", "4月", "5月", "6月"], WIDE)
    panel(s, 696, 400, 520, 128, BG_KEY, None)
    t(s, 720, 416, 472, 96,
      "行＝店舗（store）\n列＝月（month）\n交点＝売上の合計（SUM(amount)）", "Std-17N-170", TERTIARY, BG_KEY)
    t(s, 696, 552, 520, 72,
      "1店舗＝1行。月ごとの値を横に読み比べられる\n帳票・グラフ・表計算への受け渡しに向く", bullet="・")
tl.click("横持ちの表を表示（東京4月は2行が合算されて150）", g.shapes)
finish(tl)

# ================================================================ 04 全体像
s = new("変換は4つの手順で考える", "02 ── 手順の全体像")
tl = Timeline(s)
steps = [("行キー・列キー・値を決める",
          "何を行に、何を列に、何を交点に置くかを決める。列にする値は事前に列挙しておく。",
          "store, month, amount"),
         ("CASE式で列ごとに振り分ける",
          "列にしたい値ごとにCASE式を1つ書き、条件に合う行の値だけを取り出す。",
          "CASE WHEN … END"),
         ("GROUP BYで1行に畳む",
          "行キーでグループ化し、SUMなどの集計関数で各列を1つの値にまとめる。",
          "SUM(…) GROUP BY"),
         ("NULL・合計・並び順を整える",
          "データのない交点の表示を決め、合計列や ORDER BY を加えて読みやすくする。",
          "COALESCE(…, 0)")]
gap = 40
cw = (W - gap * 3) / 4
for i, (head, body, snippet) in enumerate(steps):
    x = MARGIN_X + i * (cw + gap)
    with grab(s) as g:
        if i:
            arrow_right(s, x - gap / 2, 364, 20, BORDER)
        panel(s, x, 200, cw, 328, WHITE, BORDER, RADIUS["none"])
        rect(s, x, 200, cw, 48, fill=PRIMARY)
        t(s, x + 24, 200, cw - 48, 48, f"手順{i + 1}", "Std-20B-150", WHITE, PRIMARY, anchor="m")
        t(s, x + 24, 272, cw - 48, 72, head, "Std-22B-150")
        t(s, x + 24, 360, cw - 48, 90, body, "Std-16N-170", INK_SUB)
        panel(s, x + 16, 464, cw - 32, 48, CODE_BG, None)
        code_line(s, x + 32, 464, cw - 48, 48, snippet)
    tl.click(f"手順{i + 1}：{head}", g.shapes)
with grab(s) as g:
    band(s, "基本形：SELECT 行キー, SUM(CASE WHEN 列キー = '値' THEN 値 END) … GROUP BY 行キー",
         y=552, h=72)
tl.click("SQLの基本形を表示", g.shapes)
finish(tl)

# ================================================================ 05 手順1
s = new("手順1｜行キー・列キー・値を決める", "02 ── 手順1",
        "縦持ちの列を、結果の「行」「列」「交点」のどれに使うかに割り当てる。")
tl = Timeline(s)
gt = grid(s, MARGIN_X, 244, [180, 180, 200], ["store", "month", "amount"], LONG, num_from=2)
roles = [("行キー", "store（店舗）", "結果の1行を決める列。GROUP BY に書く。\n東京・大阪の2行になる。"),
         ("列キー", "month（月）", "横に並べる値。4月・5月・6月がそれぞれ列になる。"),
         ("値", "amount（売上）", "交点に入れる数値。SUM などの集計関数でまとめる。")]
for i, (role, name, desc) in enumerate(roles):
    cx, cwid = gt["xs"][i], gt["widths"][i]
    y = 200 + i * 120
    with grab(s) as g:
        badge(s, cx + (cwid - 88) / 2, 204, role, fill=PRIMARY, fg=WHITE, w=88)
        panel(s, cx + 4, 240, cwid - 8, gt["bottom"] - 236, None, PRIMARY, line_px=2)
        badge(s, 672, y + 4, role, fill=PRIMARY, fg=WHITE, w=88)
        t(s, 784, y, 432, 32, name, "Std-20B-150")
        t(s, 784, y + 36, 432, 64, desc, "Std-16N-170", INK_SUB)
    tl.click(f"{role}＝{name} を示す", g.shapes)
with grab(s) as g:
    band(s, "列キーの値（4月・5月・6月）は、SQLを書く時点で決まっている必要がある", y=568)
tl.click("列は事前に固定される点を強調", g.shapes)
finish(tl)

# ================================================================ 06 手順2
s = new("手順2｜CASE式で値を列に振り分ける", "02 ── 手順2",
        "月ごとにCASE式を1つ書く。条件に合う行だけ amount を取り出す。")
tl = Timeline(s)
code = [
    "SELECT",
    "  store,",
    "  CASE WHEN month = '4月' THEN amount END AS \"4月\",",
    "  CASE WHEN month = '5月' THEN amount END AS \"5月\",",
    "  CASE WHEN month = '6月' THEN amount END AS \"6月\"",
    "FROM sales;",
]
cb = code_block(s, MARGIN_X, 200, 600, code, highlights=[(2, 2), (3, 3), (4, 4)])
gt = grid(s, 696, 200, [130, 130, 130, 130], ["store", "4月", "5月", "6月"],
          [spread(r) for r in LONG])
notes6 = ["条件に合う行は amount、合わない行は NULL（ELSE を省略した場合）",
          "行数は元のまま6行。まだ「横に広げただけ」"]
for i, m in enumerate(MONTHS):
    shapes = cb["highlights"][i] + col(gt["cells"], i + 1)
    if i == 0:
        with grab(s) as g:
            t(s, MARGIN_X, 456, 600, 30, notes6[0], "Std-16N-170", INK, bullet="・")
        shapes += g.shapes
    tl.click(f"{m}のCASE式と{m}列の値を表示", shapes)
with grab(s) as g:
    t(s, MARGIN_X, 490, 600, 30, notes6[1], "Std-16N-170", INK, bullet="・")
    band(s, "同じ店舗の行が残っている。次の手順で1行にまとめる", y=568)
tl.click("まだ6行のままである点を示す", g.shapes)
finish(tl)

# ================================================================ 07 手順3
s = new("手順3｜GROUP BY と SUM で1行にまとめる", "02 ── 手順3",
        "CASE式を集計関数で包み、行キーでグループ化する。")
tl = Timeline(s)
code = [
    "SELECT",
    "  store,",
    "  SUM(CASE WHEN month = '4月' THEN amount END) AS \"4月\",",
    "  SUM(CASE WHEN month = '5月' THEN amount END) AS \"5月\",",
    "  SUM(CASE WHEN month = '6月' THEN amount END) AS \"6月\"",
    "FROM sales",
    "GROUP BY store;",
]
cb = code_block(s, MARGIN_X, 200, 680, code, highlights=[(2, 4), (6, 6)])
with grab(s) as g1:
    t(s, MARGIN_X, 488, 680, 30, "SUM は NULL を無視して、同じグループの値を合計する", "Std-16N-170",
      bullet="・")
with grab(s) as g2:
    t(s, MARGIN_X, 520, 680, 30, "GROUP BY store で、店舗ごとに1行へ畳まれる", "Std-16N-170",
      bullet="・")
tl.click("SUM で包んだ部分を示す", cb["highlights"][0] + g1.shapes)
tl.click("GROUP BY を示す", cb["highlights"][1] + g2.shapes)
t(s, 776, 200, 440, 28, "実行結果（2行）", "Std-18B-160")
gt = grid(s, 776, 236, [110, 110, 110, 110], ["store", "4月", "5月", "6月"], WIDE)
tl.click("実行結果の行を表示", row(gt["cells"], 0))
tl.then(row(gt["cells"], 1))
with grab(s) as g:
    panel(s, 776, 376, 440, 80, BG_KEY, None)
    t(s, 800, 376, 392, 80, "東京の4月は2行ある\n→ 120 + 30 = 150 に合算", "Std-17N-170", TERTIARY, BG_KEY,
      anchor="m")
tl.click("東京4月＝150 の理由", g.shapes)
with grab(s) as g:
    panel(s, 776, 472, 440, 80, BG_KEY, None)
    t(s, 800, 472, 392, 80, "大阪の6月は該当する行がない\n→ NULL のまま（手順4で扱う）", "Std-17N-170",
      TERTIARY, BG_KEY, anchor="m")
tl.click("大阪6月＝NULL の理由", g.shapes)
with grab(s) as g:
    band(s, "6行が2行になり、横持ちの形になった", y=576, h=56)
tl.then(g.shapes, gap=300)
finish(tl)

# ================================================================ 08 流れを通しで見る
s = new("6行が2行になるまでを通しで見る", "02 ── 手順2〜3のまとめ")
tl = Timeline(s)
t(s, MARGIN_X, 196, 340, 28, "元データ（6行）", "Std-18B-160")
t(s, 450, 196, 360, 28, "CASE式の結果（6行）", "Std-18B-160")
t(s, 856, 196, 360, 28, "GROUP BY の結果（2行）", "Std-18B-160")
grid(s, MARGIN_X, 232, [100, 100, 140], ["store", "month", "amount"], LONG, num_from=2)
arrow_right(s, 427, 356, 20, BORDER)
arrow_right(s, 833, 356, 20, BORDER)
ct = grid(s, 450, 232, [90, 90, 90, 90], ["store", "4月", "5月", "6月"], [spread(r) for r in LONG])
tl.click("CASE式で1行ずつ横に広げる（自動で6行続く）", row(ct["cells"], 0))
for r in range(1, 6):
    tl.then(row(ct["cells"], r), gap=80)
groups = [("東京", 0, 4), ("大阪", 4, 2)]
outlines = []
for name, start, n in groups:
    with grab(s) as g:
        panel(s, 452, ct["top"] + start * 41 + 2, 356, n * 41 - 4, None, PRIMARY, line_px=2)
    outlines.append(g.shapes)
tl.click("store ごとのグループを囲む", outlines[0])
tl.then(outlines[1])
rt = grid(s, 856, 232, [90, 90, 90, 90], ["store", "4月", "5月", "6月"], WIDE)
with grab(s) as g0:
    t(s, 856, 372, 360, 30, "東京：4行 → 1行（4月は 120 + 30）", "Std-16N-170", INK_SUB, bullet="・")
with grab(s) as g1:
    t(s, 856, 404, 360, 30, "大阪：2行 → 1行（6月は NULL）", "Std-16N-170", INK_SUB, bullet="・")
tl.click("SUM で畳んだ結果を表示", row(rt["cells"], 0) + g0.shapes)
tl.then(row(rt["cells"], 1) + g1.shapes)
with grab(s) as g:
    band(s, "CASE式で横に広げ、GROUP BY で縦に畳む ─ これがクロス集計の基本形", y=568)
tl.click("まとめの一文を表示", g.shapes)
finish(tl)

# ================================================================ 09 手順4
s = new("手順4｜NULL・合計・並び順を整える", "02 ── 手順4")
tl = Timeline(s)
code = [
    "SELECT",
    "  store,",
    "  COALESCE(SUM(CASE WHEN month = '4月' THEN amount END), 0) AS \"4月\",",
    "  COALESCE(SUM(CASE WHEN month = '5月' THEN amount END), 0) AS \"5月\",",
    "  COALESCE(SUM(CASE WHEN month = '6月' THEN amount END), 0) AS \"6月\",",
    "  SUM(amount) AS 合計",
    "FROM sales",
    "GROUP BY store",
    "ORDER BY store;",
]
cb = code_block(s, MARGIN_X, 172, W, code, lh=28, pad=16, highlights=[(2, 4), (5, 5), (8, 8)])
anns = [("1", "COALESCE で NULL を 0 に置き換える", "大阪の6月が 0 になる"),
        ("2", "SUM(amount) で店舗ごとの合計列を足す", "CASE式を通さない全月の合計"),
        ("3", "ORDER BY で行の並び順を固定する", "GROUP BY だけでは順序は保証されない")]
for i, (n, head, sub) in enumerate(anns):
    y = 488 + i * 56
    with grab(s) as g:
        num(s, 720, y + 4, n, d=40)
        t(s, 776, y - 2, 440, 30, head, "Std-17N-170")
        t(s, 776, y + 28, 440, 22, sub, "Dns-14N-130", INK_SUB)
    tl.click(f"{head}", cb["highlights"][i] + g.shapes)
gt = grid(s, MARGIN_X, 488, [140, 110, 110, 110, 128], ["store", "4月", "5月", "6月", "合計"],
          [["大阪", "90", "110", "0", "200"], ["東京", "150", "140", "160", "450"]])
tl.click("完成した結果を表示", row(gt["cells"], 0))
tl.then(row(gt["cells"], 1))
finish(tl)

# ================================================================ 10 DB別
s = new("DBごとの専用構文", "03 ── 応用", "専用構文は短く書けるが、DBをまたいで移植できない。")
tl = Timeline(s)
dbs = [["標準SQL（全DB共通）", "CASE式＋集計関数", "SUM(CASE WHEN month = '4月' THEN amount END)"],
       ["PostgreSQL / SQLite", "FILTER句", "SUM(amount) FILTER (WHERE month = '4月')"],
       ["SQL Server", "PIVOT演算子", "PIVOT (SUM(amount) FOR month IN ([4月], [5月], [6月]))"],
       ["Oracle Database", "PIVOT句", "PIVOT (SUM(amount) FOR month IN ('4月' AS \"4月\", …))"],
       ["MySQL", "専用構文なし", "標準SQLのCASE式で書く"]]
widths = [240, 176, W - 416]
xs = [MARGIN_X, MARGIN_X + 240, MARGIN_X + 416]
y0 = 200
panel(s, MARGIN_X, y0, W, 44, BG_KEY, None)
for x, wd, label in zip(xs, widths, ["データベース", "書き方", "4月列の例"]):
    t(s, x + 16, y0, wd - 32, 44, label, "Dns-14B-130", PRIMARY, BG_KEY, anchor="m")
rule(s, MARGIN_X, y0 + 44, W, 2, PRIMARY)
cy = y0 + 46
for i, (db, how, ex) in enumerate(dbs):
    with grab(s) as g:
        t(s, xs[0] + 16, cy, widths[0] - 32, 60, db, "Std-16B-170", anchor="m")
        t(s, xs[1] + 16, cy, widths[1] - 32, 60, how, "Std-16N-170", INK_SUB, anchor="m")
        if ex.startswith("標準SQLの"):
            t(s, xs[2] + 16, cy, widths[2] - 32, 60, ex, "Std-16N-170", INK_SUB, anchor="m")
        else:
            code_line(s, xs[2] + 16, cy + 12, widths[2] - 32, 36, ex, bg=WHITE)
    rule(s, MARGIN_X, cy + 60, W, 1, BORDER_SOFT)
    if i:
        tl.click(f"{db}：{how}", g.shapes)
    cy += 61
with grab(s) as g:
    band(s, "迷ったら CASE式で書く。移植しやすく、条件も自由に書ける", y=568)
tl.click("推奨の書き方を表示", g.shapes)
finish(tl)

# ================================================================ 11 注意点
s = new("つまずきやすいポイント", "03 ── 注意点")
tl = Timeline(s)
pits = [("列は事前に固定される",
         "列キーの値はSQLに直接書く。7月が増えたら1行足す。値が増え続けるなら動的SQLやBIツールで対応する。",
         "SUM(CASE WHEN month = '7月' THEN amount END)"),
        ("COUNT では ELSE 0 を書かない",
         "ELSE 0 だと 0 も1件と数え、全行数になる。ELSE を省いて NULL にする（または SUM で数える）。",
         "COUNT(CASE WHEN month = '4月' THEN 1 END)"),
        ("NULL と 0 を区別する",
         "「データなし」を0に置き換えてよいか確認する。AVG で平均を取る場合、0埋めすると結果が変わる。",
         "AVG(CASE WHEN month = '6月' THEN amount END)"),
        ("集計前に粒度をそろえる",
         "重複行や JOIN による行の増加があると合計が膨らむ。結合してから集計する場合は件数を確かめる。",
         "SELECT COUNT(*) FROM sales")]
cw = (W - 16) / 2
for i, (head, body, ex) in enumerate(pits):
    x = MARGIN_X + (i % 2) * (cw + 16)
    y = 172 + (i // 2) * 232
    with grab(s) as g:
        panel(s, x, y, cw, 216, WHITE, BORDER, RADIUS["none"])
        rect(s, x, y, 6, 216, fill=WARN_ORANGE)
        badge(s, x + 32, y + 24, f"注意{i + 1}", fill=WARN_ORANGE, fg=WHITE, w=72)
        t(s, x + 120, y + 20, cw - 144, 36, head, "Std-20B-150")
        t(s, x + 32, y + 68, cw - 56, 60, body, "Std-16N-170", INK_SUB)
        panel(s, x + 32, y + 144, cw - 56, 48, CODE_BG, None)
        code_line(s, x + 48, y + 144, cw - 80, 48, ex)
    tl.click(f"注意{i + 1}：{head}", g.shapes)
finish(tl)

# ================================================================ 12 まとめ
s = new("まとめ", "04 ── まとめ")
tl = Timeline(s)
recap = [("行キー・列キー・値を決める", "列にする値（4月・5月・6月）は事前に列挙する"),
         ("CASE式で列ごとに振り分ける", "条件に合う行だけ値を取り出し、合わない行は NULL"),
         ("GROUP BY と SUM で1行に畳む", "行キーごとに集計し、縦に並んだ行をまとめる"),
         ("NULL・合計・並び順を整える", "COALESCE・合計列・ORDER BY で読みやすく仕上げる")]
for i, (head, sub) in enumerate(recap):
    y = 176 + i * 92
    with grab(s) as g:
        num(s, MARGIN_X, y + 8, str(i + 1))
        t(s, MARGIN_X + 72, y, 520, 36, head, "Std-22B-150")
        t(s, MARGIN_X + 72, y + 38, 1000, 30, sub, "Std-17N-170", INK_SUB)
        if i < 3:
            rule(s, MARGIN_X + 72, y + 80, W - 72, 1, BORDER_SOFT)
    tl.click(f"まとめ{i + 1}：{head}", g.shapes)
with grab(s) as g:
    band(s, "保存は縦持ち、見せるときに横持ち。変換はSQLの最後に行う", y=560, h=72)
tl.click("原則を表示", g.shapes)
finish(tl)

# ================================================================ 13 出典
s = new("本資料について（出典・注記）", "APPENDIX")
notes = [("デザイン", "デジタル庁デザインシステムβ版の基本デザインを参考に構成しています。"),
         ("データ", "サンプルの sales テーブルと数値は説明用の架空のものです。"),
         ("SQL", "構文は代表的な書き方です。対応状況はDBのバージョンで異なるため、各DBの公式ドキュメントで確認してください。"),
         ("動き", "アニメーションはフェードのみ・クリック進行です。動きがなくても内容が伝わる構成にしています。"),
         ("書体", "Noto Sans JP（SIL Open Font License 1.1）。SQLコードも同じ書体で組んでいます。")]
y = 184
for head, body in notes:
    badge(s, MARGIN_X, y + 2, head, fill=BG_KEY, fg=PRIMARY, w=104)
    t(s, MARGIN_X + 128, y, W - 128, 32, body, "Std-16N-170")
    y += 60
rule(s, MARGIN_X, 516, W, 2, BORDER)
t(s, MARGIN_X, 540, W, 30,
  "出典：デジタル庁デザインシステムウェブサイト https://design.digital.go.jp/dads/ のコンテンツを加工して作成",
  "Std-16N-170", INK_SUB)
t(s, MARGIN_X, 576, W, 24, "本資料はデジタル庁が作成・監修したものではありません。", "Dns-14N-130", INK_MUTE)
finish(Timeline(s))

assert len(prs.slides) == TOTAL, len(prs.slides)
prs.save(PPTX)
(OUT / "SQLクロス集計_アニメーション手順.json").write_text(
    json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print("saved:", PPTX)
print("clicks:", sum(len(m["clicks"]) for m in manifest))
